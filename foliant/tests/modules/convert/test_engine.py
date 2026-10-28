"""Test del motore di conversione da PDF (CNV-01, 02, 03, 07, 09 e formati P1/P2)."""

from __future__ import annotations

import csv
import io
import os
import zipfile

import pymupdf
import pytest

from foliant.api import OperationCancelled
from foliant.modules.convert.common import ConversionError, RangeError, parse_page_range
from foliant.modules.convert.engine import convert_pdf, run_conversion


# ------------------------------------------------------------------ fixture
def _save(doc, path) -> str:
    doc.save(str(path))
    doc.close()
    return str(path)


@pytest.fixture
def two_columns(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 60), "Titolo del documento", fontsize=18)
    left = pymupdf.Rect(50, 100, 280, 800)
    right = pymupdf.Rect(315, 100, 545, 800)
    page.insert_textbox(left, "SINISTRA uno. " + "alfa beta gamma delta " * 30 + "FINESINISTRA", fontsize=10)
    page.insert_textbox(right, "DESTRA uno. " + "epsilon zeta eta theta " * 30 + "FINEDESTRA", fontsize=10)
    return _save(doc, tmp_path / "colonne.pdf")


@pytest.fixture
def links_pdf(tmp_path):
    doc = pymupdf.open()
    for i in range(3):
        doc.new_page(width=595, height=842)
    p0 = doc[0]
    for k, url in enumerate(["https://a.example.com/", "https://b.example.com/x", "https://c.example.com/y"]):
        y = 100 + k * 40
        p0.insert_text((72, y), f"Collegamento numero {k + 1}", fontsize=12)
        p0.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(70, y - 12, 250, y + 4), "uri": url})
    for k, target in enumerate([1, 2]):
        y = 300 + k * 40
        p0.insert_text((72, y), f"Vai alla pagina {target + 1}", fontsize=12)
        p0.insert_link({"kind": pymupdf.LINK_GOTO, "from": pymupdf.Rect(70, y - 12, 250, y + 4), "page": target,
                        "to": pymupdf.Point(72, 100)})
    doc[1].insert_text((72, 100), "Destinazione due", fontsize=12)
    doc[2].insert_text((72, 100), "Destinazione tre", fontsize=12)
    return _save(doc, tmp_path / "link.pdf")


def _jpeg_bytes(w=320, h=200):
    from PIL import Image
    img = Image.new("RGB", (w, h))
    for x in range(w):
        for y in range(0, h, 4):
            img.putpixel((x, y), (x % 256, y % 256, 120))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue()


@pytest.fixture
def image_pdf(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 80), "Didascalia della fotografia", fontsize=12)
    page.insert_image(pymupdf.Rect(72, 100, 392, 300), stream=_jpeg_bytes())
    return _save(doc, tmp_path / "immagine.pdf")


@pytest.fixture
def scanned_mixed(tmp_path):
    """3 pagine: 1–2 native, 3 scansionata (immagine di testo senza livello di testo)."""
    src = pymupdf.open()
    sp = src.new_page(width=595, height=842)
    sp.insert_text((60, 200), "SCANSIONE PROVA", fontsize=40)
    sp.insert_text((60, 300), "Orologio automatico", fontsize=40)
    pix = sp.get_pixmap(dpi=200)
    doc = pymupdf.open()
    for i in range(2):
        p = doc.new_page(width=595, height=842)
        p.insert_text((72, 100), f"Testo nativo pagina {i + 1}", fontsize=14)
    p = doc.new_page(width=595, height=842)
    p.insert_image(p.rect, stream=pix.tobytes("png"))
    return _save(doc, tmp_path / "misto.pdf")


# ------------------------------------------------------------------ CNV-01
def test_parse_page_range():
    assert parse_page_range("2-4, 10", 20) == [1, 2, 3, 9]
    assert parse_page_range("", 3) == [0, 1, 2]
    assert parse_page_range("18-fine", 20) == [17, 18, 19]
    assert parse_page_range("pari", 5) == [1, 3]
    assert parse_page_range("dispari", 5) == [0, 2, 4]
    assert parse_page_range("corrente", 5, current=3) == [3]
    assert parse_page_range("3-1", 5) == [2, 1, 0]
    for bad in ("25-30", "0", "abc", "1-x"):
        with pytest.raises(RangeError):
            parse_page_range(bad, 20)


def test_range_txt_order(twenty_pages, tmp_path):
    out = tmp_path / "out" / "venti.txt"
    paths = convert_pdf(twenty_pages, out, "txt", pages="2-4, 10")
    text = open(paths[0], encoding="utf-8").read()
    nums = [int(line.split()[-1]) for line in text.replace("\f", "\n").splitlines() if line.startswith("Pagina")]
    assert nums == [2, 3, 4, 10]


def test_invalid_range_no_file(twenty_pages, tmp_path):
    out = tmp_path / "dest"
    with pytest.raises(RangeError):
        convert_pdf(twenty_pages, out / "x.txt", "txt", pages="25-30")
    assert not out.exists() or not list(out.iterdir())


def test_cancel_leaves_nothing(twenty_pages, tmp_path):
    from foliant.api import Progress
    out = tmp_path / "dest"
    out.mkdir()
    progress = Progress()
    calls = {"n": 0}
    original = progress.update

    def update(fraction=None, text=None):
        calls["n"] += 1
        if calls["n"] == 5:
            progress.cancel()
        original(fraction, text)

    progress.update = update
    for fmt in ("txt", "docx", "png"):
        calls["n"] = 0
        progress._cancel.clear()
        with pytest.raises(OperationCancelled):
            convert_pdf(twenty_pages, out / f"venti.{fmt}", fmt, progress=progress)
        assert list(out.iterdir()) == []


def test_rename_policy_keeps_existing(sample_pdf, tmp_path):
    target = tmp_path / "esistente.txt"
    target.write_bytes(b"originale")
    paths = convert_pdf(sample_pdf, target, "txt")
    assert paths == [str(tmp_path / "esistente (1).txt")]
    assert target.read_bytes() == b"originale"
    paths = convert_pdf(sample_pdf, target, "txt", policy="overwrite")
    assert paths == [str(target)] and target.read_bytes() != b"originale"


def test_extraction_forbidden(tmp_path):
    path = tmp_path / "vietato.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "Riservato", fontsize=12)
    doc.save(str(path), encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="capo", user_pw="",
             permissions=pymupdf.PDF_PERM_PRINT)
    doc.close()
    with pytest.raises(ConversionError) as e:
        convert_pdf(str(path), tmp_path / "x.docx", "docx")
    assert "copia del contenuto" in str(e.value)
    assert not (tmp_path / "x.docx").exists()
    # con la password del proprietario la conversione è permessa
    assert convert_pdf(str(path), tmp_path / "x.txt", "txt", password="capo")


def test_encrypted_with_password(encrypted_pdf, tmp_path):
    from foliant.api import PasswordRequired
    with pytest.raises(PasswordRequired):
        convert_pdf(encrypted_pdf, tmp_path / "p.txt", "txt")
    paths = convert_pdf(encrypted_pdf, tmp_path / "p.txt", "txt", password="segreta")
    assert "Riservato" in open(paths[0], encoding="utf-8").read()


def test_bytes_source_and_signed_warning(sample_pdf, tmp_path):
    data = open(sample_pdf, "rb").read()
    result = run_conversion(data, tmp_path / "da_byte.txt", "txt")
    assert "Pagina 1" in open(result.paths[0], encoding="utf-8").read()


# ------------------------------------------------------------------ CNV-09
def test_txt_two_columns_reading_order(two_columns, tmp_path):
    text = open(convert_pdf(two_columns, tmp_path / "c.txt", "txt")[0], encoding="utf-8").read()
    assert text.index("Titolo") < text.index("SINISTRA")
    assert text.index("FINESINISTRA") < text.index("DESTRA uno")


def test_txt_ligature_and_form_feed(tmp_path):
    doc = pymupdf.open()
    for i in range(3):
        p = doc.new_page()
        p.insert_text((72, 100), "la ﬁne del contratto" if i == 0 else f"pagina {i + 1}", fontsize=12,
                      fontname="dejavu", fontfile="/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    path = _save(doc, tmp_path / "lig.pdf")
    out = convert_pdf(path, tmp_path / "lig.txt", "txt", page_separator="ff")[0]
    raw = open(out, "rb").read()
    text = raw.decode("utf-8")
    assert "fine" in text and "ﬁ" not in text
    assert text.count("\f") == 2


def test_txt_layout_and_encoding(sample_pdf, tmp_path):
    out = convert_pdf(sample_pdf, tmp_path / "l.txt", "txt", mode="layout", encoding="utf-16", line_ending="crlf")[0]
    text = open(out, encoding="utf-16").read()
    assert "Pagina 1" in text and "\r\n" in open(out, "rb").read().decode("utf-16")


# ------------------------------------------------------------------ CNV-03
def _docx_text(path):
    import docx
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs), d


def test_docx_flow_text_and_order(two_columns, tmp_path):
    text, _d = _docx_text(convert_pdf(two_columns, tmp_path / "c.docx", "docx", layout="flow")[0])
    assert text.index("FINESINISTRA") < text.index("DESTRA uno")
    with pymupdf.open(two_columns) as doc:
        src_words = doc[0].get_text().split()
    out_words = text.split()
    common = sum(1 for w in src_words if w in set(out_words))
    assert common / len(src_words) >= 0.99


def test_docx_table(table_pdf, tmp_path):
    path, rows = table_pdf
    _t, d = _docx_text(convert_pdf(path, tmp_path / "t.docx", "docx")[0])
    assert len(d.tables) == 1
    t = d.tables[0]
    assert (len(t.rows), len(t.columns)) == (5, 4)
    assert [[c.text for c in r.cells] for r in t.rows] == rows


def test_docx_links(links_pdf, tmp_path):
    out = convert_pdf(links_pdf, tmp_path / "l.docx", "docx")[0]
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode()
        rels = z.read("word/_rels/document.xml.rels").decode()
    for url in ("https://a.example.com/", "https://b.example.com/x", "https://c.example.com/y"):
        assert url in rels
    assert xml.count("<w:hyperlink") == 5
    assert xml.count('w:anchor="page_2"') == 1 and xml.count('w:anchor="page_3"') == 1
    assert 'w:name="page_2"' in xml and 'w:name="page_3"' in xml


def test_docx_image_original_pixels(image_pdf, tmp_path):
    out = convert_pdf(image_pdf, tmp_path / "i.docx", "docx")[0]
    from PIL import Image
    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        assert len(media) == 1
        img = Image.open(io.BytesIO(z.read(media[0])))
    assert img.size == (320, 200)


def test_docx_faithful_sections(tmp_path):
    doc = pymupdf.open()
    for i in range(6):
        p = doc.new_page(width=595.28, height=841.89)
        p.insert_text((57, 100), f"Pagina fedele {i + 1}", fontsize=14)
    path = _save(doc, tmp_path / "sei.pdf")
    _t, d = _docx_text(convert_pdf(path, tmp_path / "sei.docx", "docx", layout="faithful")[0])
    for s in d.sections:
        assert abs(s.page_width.mm - 210) <= 1 and abs(s.page_height.mm - 297) <= 1
    with zipfile.ZipFile(tmp_path / "sei.docx") as z:
        xml = z.read("word/document.xml").decode()
    breaks = xml.count("<w:sectPr") - 1 + xml.count('w:type="page"')
    assert breaks >= 5


# ------------------------------------------------------------------ CNV-02
def test_ocr_only_scanned_page(scanned_mixed, tmp_path):
    from foliant.modules.convert.extract import ocr_available
    if not ocr_available():
        pytest.skip("tesseract non installato")
    result = run_conversion(scanned_mixed, tmp_path / "m.txt", "txt")
    assert result.ocr_pages == [3]
    pages = open(result.paths[0], encoding="utf-8").read().split("\f")
    with pymupdf.open(scanned_mixed) as doc:
        for i in range(2):
            assert pages[i].strip() == doc[i].get_text().strip()
    assert "SCANSIONE" in pages[2].upper() and "OROLOGIO" in pages[2].upper()


def test_ocr_never_warns(scanned_mixed, tmp_path):
    result = run_conversion(scanned_mixed, tmp_path / "m.txt", "txt", ocr="never")
    assert result.scanned_pages == [3]
    assert any("Pagina 3" in w for w in result.warnings)


# ------------------------------------------------------------------ CNV-07
def test_png_300dpi(sample_pdf, tmp_path):
    from PIL import Image
    doc = pymupdf.open()
    for i in range(3):
        doc.new_page(width=595.28, height=841.89).insert_text((72, 100), f"A4 {i}", fontsize=12)
    path = _save(doc, tmp_path / "doc.pdf")
    paths = convert_pdf(path, tmp_path / "img" / "doc.png", "png", dpi=300)
    assert [os.path.basename(p) for p in paths] == ["doc_001.png", "doc_002.png", "doc_003.png"]
    for p in paths:
        img = Image.open(p)
        assert abs(img.size[0] - 2480) <= 1 and abs(img.size[1] - 3508) <= 1
        assert round(img.info["dpi"][0]) == 300


def test_tiff_multipage_g4(sample_pdf, tmp_path):
    from PIL import Image
    paths = convert_pdf(sample_pdf, tmp_path / "doc.tif", "tiff", multipage=True, color="bw",
                        tiff_compression="group4", dpi=100)
    assert len(paths) == 1
    img = Image.open(paths[0])
    assert img.n_frames == 3
    for k in range(3):
        img.seek(k)
        assert img.mode == "1"
        assert img.info.get("compression") == "group4"


def test_gif_palette_and_size(sample_pdf, tmp_path):
    from PIL import Image
    gif = convert_pdf(sample_pdf, tmp_path / "g.gif", "gif", pages="1", dpi=100)[0]
    png = convert_pdf(sample_pdf, tmp_path / "g.png", "png", pages="1", dpi=100)[0]
    g = Image.open(gif)
    assert g.mode == "P" and len(g.getpalette()) // 3 <= 256
    assert g.size == Image.open(png).size


def test_rotated_page_and_jpeg_bmp(tmp_path):
    from PIL import Image
    doc = pymupdf.open()
    p = doc.new_page(width=595, height=842)
    p.insert_text((72, 100), "Ruotata", fontsize=12)
    p.set_rotation(90)
    path = _save(doc, tmp_path / "rot.pdf")
    for fmt in ("jpeg", "bmp"):
        out = convert_pdf(path, tmp_path / f"r.{fmt}", fmt, dpi=72)[0]
        w, h = Image.open(out).size
        assert w > h and abs(w - 842) <= 1 and abs(h - 595) <= 1


def test_annotations_excluded(tmp_path):
    from PIL import Image, ImageChops
    doc = pymupdf.open()
    p = doc.new_page(width=300, height=300)
    p.insert_text((20, 50), "Testo", fontsize=12)
    path_plain = _save(doc, tmp_path / "plain.pdf")
    doc = pymupdf.open(path_plain)
    doc[0].add_rect_annot(pymupdf.Rect(50, 50, 250, 250)).update()
    path_ann = _save(doc, tmp_path / "ann.pdf")
    without = convert_pdf(path_ann, tmp_path / "a.png", "png", dpi=100, annotations=False)[0]
    plain = convert_pdf(path_plain, tmp_path / "p.png", "png", dpi=100)[0]
    with_ann = convert_pdf(path_ann, tmp_path / "w.png", "png", dpi=100, annotations=True)[0]
    a, b = Image.open(without).convert("RGB"), Image.open(plain).convert("RGB")
    assert ImageChops.difference(a, b).getbbox() is None
    assert ImageChops.difference(Image.open(with_ann).convert("RGB"), b).getbbox() is not None


def test_image_size_limit(tmp_path):
    doc = pymupdf.open()
    doc.new_page(width=14000, height=200)
    path = _save(doc, tmp_path / "largo.pdf")
    with pytest.raises(ConversionError) as e:
        convert_pdf(path, tmp_path / "x.jpg", "jpeg", dpi=600)
    assert "dpi" in str(e.value)


def test_long_image(tmp_path):
    from PIL import Image
    doc = pymupdf.open()
    for i in range(5):
        doc.new_page(width=595.28, height=841.89).insert_text((72, 100), f"P{i}", fontsize=12)
    path = _save(doc, tmp_path / "cinque.pdf")
    out = convert_pdf(path, tmp_path / "lunga.png", "long-image", dpi=96, spacing=0)
    assert len(out) == 1
    w, h = Image.open(out[0]).size
    assert abs(w - 794) <= 5 and abs(h - 5615) <= 5
    parts = convert_pdf(path, tmp_path / "lunga.jpg", "long-image", dpi=96, max_height=2500)
    assert len(parts) == 3 and all("_parte" in p for p in parts)
    sizes = [Image.open(p).size[1] for p in parts]
    assert all(s <= 2500 for s in sizes) and abs(sum(sizes) - h) <= 5


# ------------------------------------------------------------- CNV-04/05
@pytest.fixture
def amounts_pdf(tmp_path):
    """Tabella 20 righe × 6 colonne con importi «1.234,56» e riga dei totali."""
    import random
    random.seed(4)
    header = ["Codice", "Articolo", "Q1", "Q2", "Q3", "Q4"]
    rows, totals = [], [0.0] * 4
    for r in range(18):
        vals = [round(random.uniform(10, 9999), 2) for _ in range(4)]
        totals = [t + v for t, v in zip(totals, vals)]
        rows.append([f"00{r + 100}", f"Orologio {r}"] + [f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                                                        for v in vals])
    rows.append(["", "Totale"] + [f"{t:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") for t in totals])
    data = [header] + rows
    doc = pymupdf.open()
    page = doc.new_page(width=842, height=595)
    x0, y0, cw, rh = 40, 40, 120, 24
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            rect = pymupdf.Rect(x0 + c * cw, y0 + r * rh, x0 + (c + 1) * cw, y0 + (r + 1) * rh)
            page.draw_rect(rect, color=(0, 0, 0), width=0.6)
            if val:
                page.insert_text((rect.x0 + 3, rect.y1 - 7), val, fontsize=10)
    return _save(doc, tmp_path / "importi.pdf"), totals


def test_xlsx_numbers_and_totals(amounts_pdf, tmp_path):
    from openpyxl import load_workbook
    path, totals = amounts_pdf
    out = convert_pdf(path, tmp_path / "i.xlsx", "xlsx")[0]
    ws = load_workbook(out).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0] == ["Codice", "Articolo", "Q1", "Q2", "Q3", "Q4"]
    body, total_row = values[1:-1], values[-1]
    assert ws.cell(row=2, column=1).value == "00100" and ws.cell(row=2, column=1).data_type == "s"
    for c in range(4):
        col = [r[2 + c] for r in body]
        assert all(isinstance(v, float) for v in col)
        assert abs(sum(col) - total_row[2 + c]) < 0.01
        assert abs(total_row[2 + c] - round(totals[c], 2)) < 0.01


def test_xlsx_merge_pages(tmp_path):
    from openpyxl import load_workbook
    doc = pymupdf.open()
    n = 0
    for pg in range(3):
        page = doc.new_page(width=595, height=842)
        data = [["Data", "Descrizione", "Importo"]] + [[f"0{k + 1}/01/2025", f"Voce {n + k}", f"{n + k},00"]
                                                       for k in range(5)]
        n += 5
        for r, row in enumerate(data):
            for c, val in enumerate(row):
                rect = pymupdf.Rect(50 + c * 150, 50 + r * 22, 200 + c * 150, 72 + r * 22)
                page.draw_rect(rect, color=(0, 0, 0), width=0.6)
                page.insert_text((rect.x0 + 3, rect.y1 - 6), val, fontsize=10)
    path = _save(doc, tmp_path / "lunga.pdf")
    ws = load_workbook(convert_pdf(path, tmp_path / "m.xlsx", "xlsx", merge_tables=True)[0]).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ["Data", "Descrizione", "Importo"]
    assert sum(1 for r in rows if r[0] == "Data") == 1
    assert [r[1] for r in rows[1:]] == [f"Voce {k}" for k in range(15)]
    import datetime
    assert isinstance(rows[1][0], datetime.datetime)


def test_xlsx_two_level_header(tmp_path):
    from openpyxl import load_workbook
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    # intestazione superiore unita su 2 colonne + 2 colonne
    page.draw_rect(pymupdf.Rect(50, 50, 250, 72), color=(0, 0, 0))
    page.insert_text((55, 66), "Vendite", fontsize=10)
    page.draw_rect(pymupdf.Rect(250, 50, 450, 72), color=(0, 0, 0))
    page.insert_text((255, 66), "Acquisti", fontsize=10)
    for r in range(1, 4):
        for c in range(4):
            rect = pymupdf.Rect(50 + c * 100, 50 + r * 22, 150 + c * 100, 72 + r * 22)
            page.draw_rect(rect, color=(0, 0, 0))
            page.insert_text((rect.x0 + 3, rect.y1 - 6), ["Q", "V", "Q", "V"][c] if r == 1 else str(r * 10 + c),
                             fontsize=10)
    path = _save(doc, tmp_path / "h2.pdf")
    ws = load_workbook(convert_pdf(path, tmp_path / "h2.xlsx", "xlsx")[0]).active
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert "A1:B1" in merged and "C1:D1" in merged


def test_csv_rfc4180_and_matrix(tmp_path):
    from openpyxl import load_workbook
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    data = [["Nome", "Nota"], ['Rossi; "Mario"', "riga uno"], ["Bianchi", "ok"]]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            rect = pymupdf.Rect(50 + c * 200, 50 + r * 40, 250 + c * 200, 90 + r * 40)
            page.draw_rect(rect, color=(0, 0, 0))
            lines = val.split("\n")
            for k, line in enumerate(lines):
                page.insert_text((rect.x0 + 3, rect.y0 + 14 + k * 12), line, fontsize=10)
    page.insert_text((253, 50 + 40 + 26), "e due", fontsize=10)   # seconda riga nella cella «Nota»
    path = _save(doc, tmp_path / "csv.pdf")
    out = convert_pdf(path, tmp_path / "t.csv", "csv", delimiter=";")
    assert [os.path.basename(p) for p in out] == ["t_p1_t1.csv"]
    with open(out[0], encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert rows[1][0] == 'Rossi; "Mario"'
    assert rows[1][1] == "riga uno\ne due"
    ws = load_workbook(convert_pdf(path, tmp_path / "t.xlsx", "xlsx")[0]).active
    matrix = [[("" if c.value is None else str(c.value)) for c in r] for r in ws.iter_rows()]
    assert matrix == rows


def test_csv_no_tables_no_file(sample_pdf, tmp_path):
    with pytest.raises(ConversionError) as e:
        convert_pdf(sample_pdf, tmp_path / "x.csv", "csv")
    assert "tabella" in str(e.value)
    assert list(tmp_path.glob("*.csv")) == []


# ------------------------------------------------------------------ CNV-06
def test_pptx_slides(tmp_path):
    from pptx import Presentation
    doc = pymupdf.open()
    for i in range(12):
        p = doc.new_page(width=960, height=540)
        p.insert_text((60, 80), f"Diapositiva {i + 1}", fontsize=28)
        p.insert_text((60, 140), "Collezione autunno orologi", fontsize=16)
    path = _save(doc, tmp_path / "pres.pdf")
    prs = Presentation(convert_pdf(path, tmp_path / "p.pptx", "pptx")[0])
    assert len(prs.slides) == 12
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    for i, slide in enumerate(prs.slides):
        text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        assert f"Diapositiva {i + 1}" in text and "Collezione autunno orologi" in text


def test_pptx_image_and_a4(image_pdf, tmp_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(convert_pdf(image_pdf, tmp_path / "i.pptx", "pptx")[0])
    assert abs(prs.slide_width / prs.slide_height - 210 / 297) / (210 / 297) < 0.005
    shapes = list(prs.slides[0].shapes)
    pics = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    texts = [s for s in shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert len(pics) == 1 and texts
    pic = pics[0]
    k = prs.slide_width / 595
    assert abs(pic.left - 72 * k) <= 0.02 * prs.slide_width
    assert abs(pic.top - 100 * k) <= 0.02 * prs.slide_width


# ------------------------------------------------------------------ CNV-11
def test_html_flow_headings_and_single_file(tmp_path):
    doc = pymupdf.open()
    for i in range(3):
        p = doc.new_page()
        p.insert_text((72, 80), f"Capitolo {i + 1}", fontsize=24)
        p.insert_text((72, 120), "Testo del paragrafo con < e & speciali", fontsize=11)
        if i == 0:
            p.insert_image(pymupdf.Rect(72, 150, 272, 275), stream=_jpeg_bytes())
    path = _save(doc, tmp_path / "h.pdf")
    out = convert_pdf(path, tmp_path / "h.html", "html", layout="flow")[0]
    html = open(out, encoding="utf-8").read()
    assert html.count("<h1>") == 3
    for i in range(3):
        assert f"<h1>Capitolo {i + 1}</h1>" in html
    assert "&lt; e &amp;" in html
    assert 'src="data:image/' in html and 'src="http' not in html
    from html.parser import HTMLParser
    HTMLParser().feed(html)
    # cartella di risorse
    outs = convert_pdf(path, tmp_path / "h2.html", "html", single_file=False)
    assert any(p.endswith(".jpg") or p.endswith(".jpeg") for p in outs)
    assert os.path.isdir(tmp_path / "h2_files")


def test_html_faithful(sample_pdf, tmp_path):
    html = open(convert_pdf(sample_pdf, tmp_path / "f.html", "html", layout="faithful")[0], encoding="utf-8").read()
    assert html.count('id="page-') == 3 and "Pagina 2" in html


# ------------------------------------------------------------- CNV-12/13
def test_xml_words(sample_pdf, tmp_path):
    import xml.etree.ElementTree as ET
    out = convert_pdf(sample_pdf, tmp_path / "x.xml", "xml", pages="1-2")[0]
    root = ET.parse(out).getroot()
    ns = {"f": "urn:foliant:export:1"}
    pages = root.findall("f:page", ns)
    assert len(pages) == 2
    txt = open(convert_pdf(sample_pdf, tmp_path / "x.txt", "txt", pages="1-2")[0], encoding="utf-8").read()
    for page in pages:
        w, h = float(page.get("width")), float(page.get("height"))
        words = page.findall(".//f:word", ns)
        assert words
        for word in words:
            x0, y0, x1, y1 = map(float, word.get("bbox").split())
            assert 0 <= x0 <= x1 <= w and 0 <= y0 <= y1 <= h
        joined = " ".join(word.text for word in words)
        assert joined.split()[0] in txt


def test_xml_schema_published():
    from foliant.modules.convert.export_text import XSD_PATH
    assert XSD_PATH.exists()


def test_xml_annotations(sample_pdf, tmp_path):
    import xml.etree.ElementTree as ET
    doc = pymupdf.open(sample_pdf)
    for k in range(4):
        doc[0].add_text_annot((100 + k * 30, 200), f"Nota {k}")
    path = _save(doc, tmp_path / "ann.pdf")
    root = ET.parse(convert_pdf(path, tmp_path / "a.xml", "xml")[0]).getroot()
    anns = root.findall(".//{urn:foliant:export:1}annotation")
    assert len(anns) == 4 and anns[0].get("type") == "Text" and anns[2].text == "Nota 2"


def test_epub_bookmarks_and_fixed(tmp_path):
    doc = pymupdf.open()
    for i in range(20):
        p = doc.new_page()
        p.insert_text((72, 300), f"Sezione {i // 5 + 1} pagina {i + 1}", fontsize=12)
        p.insert_text((72, 800), f"Pagina {i + 1} di 20", fontsize=9)
    doc.set_toc([[1, f"Parte {k + 1}", k * 5 + 1] for k in range(4)])
    path = _save(doc, tmp_path / "libro.pdf")
    out = convert_pdf(path, tmp_path / "l.epub", "epub")[0]
    with zipfile.ZipFile(out) as z:
        assert z.namelist()[0] == "mimetype" and z.read("mimetype") == b"application/epub+zip"
        assert z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
        nav = z.read("OEBPS/nav.xhtml").decode()
        assert all(f">Parte {k + 1}<" in nav for k in range(4))
        chapters = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("OEBPS/chapter"))
        assert "di 20" not in chapters and "Sezione 4 pagina 20" in chapters
        import xml.etree.ElementTree as ET
        for n in z.namelist():
            if n.endswith((".xhtml", ".opf", ".xml")):
                ET.fromstring(z.read(n))
    fixed = convert_pdf(path, tmp_path / "f.epub", "epub", layout="fixed", dpi=50)[0]
    with zipfile.ZipFile(fixed) as z:
        pages = [n for n in z.namelist() if n.startswith("OEBPS/page") and n.endswith(".xhtml")]
        assert len(pages) == 20 and pages == sorted(pages)
        assert "pre-paginated" in z.read("OEBPS/content.opf").decode()


# ------------------------------------------------------------------ CNV-10
def test_rtf_styles_and_images(tmp_path):
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 100), "Normale àèìòù", fontsize=12)
    p.insert_text((72, 130), "Grassetto", fontsize=12, fontname="hebo")
    p.insert_text((72, 160), "Corsivo", fontsize=12, fontname="heit")
    p.insert_text((72, 190), "Greco αβγ", fontsize=12, fontname="Helvetica")
    p.insert_image(pymupdf.Rect(72, 220, 232, 320), stream=_jpeg_bytes())
    p.insert_image(pymupdf.Rect(72, 340, 172, 400), stream=_jpeg_bytes(200, 120))
    path = _save(doc, tmp_path / "r.pdf")
    rtf = open(convert_pdf(path, tmp_path / "r.rtf", "rtf")[0], encoding="ascii").read()
    assert rtf.startswith("{\\rtf1") and rtf.count("{") == rtf.count("}") + rtf.count("\\{") - rtf.count("\\}")
    import re
    assert re.search(r"\\b\\i0 Grassetto", rtf) and re.search(r"\\b0\\i Corsivo", rtf)
    assert "\\u224?" in rtf    # à
    assert rtf.count("\\pict") == 2
