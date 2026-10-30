"""Test dell'interfaccia del modulo Converti (finestre simulate con monkeypatch)."""

from __future__ import annotations

import io
import os

import pymupdf
import pytest
from PySide6.QtWidgets import QDialog, QMessageBox

from foliant.core.platform import find_tool
from foliant.modules.convert.create_dialogs import BlankDialog, ImagesDialog
from foliant.modules.convert.export_dialog import ExportDialog
from foliant.modules.convert.optimize_dialogs import CompressDialog


@pytest.fixture
def ctrl(window):
    return window.ctx.convert_controller


@pytest.fixture
def messages(window, monkeypatch):
    """Raccoglie errori e informazioni invece di mostrare finestre modali."""
    log = {"error": [], "info": [], "confirm": [], "notify": []}
    monkeypatch.setattr(window.ctx, "error", lambda t, m: log["error"].append((t, m)))
    monkeypatch.setattr(window.ctx, "info", lambda t, m: log["info"].append((t, m)))
    monkeypatch.setattr(window.ctx, "notify", lambda text, actions=(), duration_ms=None:
                        log["notify"].append((text, [a[0] for a in actions])))

    def confirm(t, m):
        log["confirm"].append((t, m))
        return True
    monkeypatch.setattr(window.ctx, "confirm", confirm)
    monkeypatch.setattr(QMessageBox, "exec", lambda self: 0)
    return log


def _accept_export(monkeypatch, folder, fmt=None, pages="", **fields):
    def fake_exec(dlg):
        if fmt:
            dlg.format_box.setCurrentIndex(dlg.format_box.findData(fmt))
        dlg.folder_edit.setText(str(folder))
        dlg.range_edit.setText(pages)
        dlg.panels[dlg.fmt].load(fields)
        if not dlg.validate():
            return QDialog.DialogCode.Rejected
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ExportDialog, "exec", fake_exec)


# ------------------------------------------------------------- registrazione
def test_mode_and_home_actions(window, ctrl):
    assert "converti" in window.modes
    titles = [a[1] for a in window.home._actions]
    assert "Crea PDF da file" in titles and "Immagini in PDF" in titles
    mode = window.modes["converti"]
    texts = [a.text() for a in mode.toolbar.actions()]
    assert "Esporta come…" in texts and "Comprimi" in texts


def test_ribbon_groups_like_excel(window, ctrl):
    """Scheda Converti in gruppi con nome: Esporta | Crea (Crea PDF ▾) | Ottimizza (Comprimi, Altri strumenti ▾)."""
    bar = window.modes["converti"].toolbar
    groups = bar.groups()
    assert [title for title, _acts in groups] == ["Esporta", "Crea", "Ottimizza"]
    export, create, optimize = (acts for _title, acts in groups)
    acts = ctrl.actions
    assert export[0] is acts["export"] and bar.group_title(acts["export"]) == "Esporta"
    # i formati: nome intero per menu e ricerca dei comandi, nome breve nella barra
    assert [a.text() for a in export[1:]] == ["In Word", "In Excel", "In PowerPoint", "In immagini", "In testo",
                                              "In HTML", "In PDF/A"]
    assert [a.iconText() for a in export[1:]] == ["Word", "Excel", "PowerPoint", "Immagini", "Testo", "HTML",
                                                  "PDF/A"]
    assert [bar.widgetForAction(a) for a in create] == [ctrl.create_button]
    assert optimize[0] is acts["compress"] and bar.group_title(acts["compress"]) == "Ottimizza"
    assert [bar.widgetForAction(a) for a in optimize[1:]] == [ctrl.more_button]
    assert ctrl.more_button.text() == "Altri strumenti"
    create_items = ctrl.create_button.menu().actions()
    more_items = ctrl.more_button.menu().actions()
    for key in ("from_file", "from_images", "from_url", "clipboard", "scanner", "camera", "photos", "screen",
                "blank", "batch_create"):
        assert acts[key] in create_items
    for key in ("advanced", "flatten", "compress_many", "repair", "batch", "batch_create"):
        assert acts[key] in more_items
    # i pulsanti ▾ funzionano anche senza documento aperto
    assert window.current_view() is None and ctrl.create_button.isEnabled() and ctrl.more_button.isEnabled()


def test_file_tab_positions(window, ctrl):
    """Voci dell'ex menu ☰ nella scheda File: Esporta › Esporta come…, Nuovo › Crea; niente doppioni."""
    b = window.backstage
    acts = ctrl.actions
    assert b.page_of(acts["export"]) == "esporta"
    entry = next(e for e in b.entries("esporta") if e.action is acts["export"])
    assert (entry.group, entry.description) == ("Esporta", "Word, Excel, PowerPoint, immagini, testo, HTML, PDF/A")
    new = {e.action: e for e in b.entries("nuovo")}
    expected = {"blank": "Scegli formato, orientamento e numero di pagine",
                "from_file": "Da Word, Excel, immagini e altri formati",
                "batch_create": "Converte in PDF una cartella intera"}
    for key, description in expected.items():
        assert b.page_of(acts[key]) == "nuovo"
        assert (new[acts[key]].group, new[acts[key]].description) == ("Crea", description)
    # stesso gruppo «Crea» del nucleo, dopo «Nuovo documento»
    crea = [e.action for e in b.entries("nuovo") if e.group == "Crea"]
    assert crea[0] is window.act_new
    positions = [crea.index(acts[key]) for key in ("blank", "from_file", "batch_create")]
    assert positions == sorted(positions)
    assert acts["from_file"].text() == "PDF da file…" and acts["from_file"].iconText() == "Da file…"
    # «Comprimi più file…» resta solo in Altri strumenti ▾; nessuna voce in File › Altri comandi
    assert b.page_of(acts["compress_many"]) is None
    assert all(b.page_of(a) != "altro" for a in acts.values())
    window.show_backstage("nuovo")
    tiles = {t.action: t for t in b.tiles}
    assert tiles[acts["from_file"]].title_label.text() == "PDF da file…"
    assert tiles[acts["blank"]].isEnabled()                 # senza documento aperto
    window.hide_backstage()


def test_module_no_longer_uses_the_old_menu():
    from pathlib import Path

    import foliant.modules.convert as module
    folder = Path(module.__file__).parent
    assert not [p.name for p in folder.glob("*.py") if "add_menu_action" in p.read_text(encoding="utf-8")]


# ------------------------------------------------------------------ CNV-01
def test_export_open_document_with_unsaved_changes(window, ctrl, sample_pdf, tmp_path, monkeypatch, messages):
    view = window.open_path(sample_pdf)
    with view.document.change("Testo") as pdf:
        pdf[1].insert_text((72, 300), "Modifica non salvata", fontsize=12)
    before = open(sample_pdf, "rb").read()
    out = tmp_path / "out"
    _accept_export(monkeypatch, out, "txt", pages="2-3")
    result = ctrl.export()
    assert result is not None and result.paths == [str(out / "esempio.txt")]
    text = open(result.paths[0], encoding="utf-8").read()
    assert "Modifica non salvata" in text and "Pagina 1" not in text and "Pagina 3" in text
    assert open(sample_pdf, "rb").read() == before
    assert view.document.is_modified


def test_export_invalid_range_creates_nothing(window, ctrl, sample_pdf, tmp_path, monkeypatch, messages):
    window.open_path(sample_pdf)
    out = tmp_path / "out"
    _accept_export(monkeypatch, out, "txt", pages="25-30")
    assert ctrl.export() is None
    assert not out.exists()
    dlg = ExportDialog(window, page_count=3, name="x", folder=str(out))
    dlg.range_edit.setText("25-30")
    assert not dlg.validate()
    assert dlg.range_error.isVisibleTo(dlg) and "3 pagine" in dlg.range_error.text()


def test_export_dialog_remembers_options(window, settings, tmp_path):
    dlg = ExportDialog(window, page_count=3, name="x", folder=str(tmp_path), fmt="png", settings=settings)
    dlg.panels["png"].fields["dpi"].setValue(300)
    dlg.accept()
    dlg2 = ExportDialog(window, page_count=3, name="x", folder=str(tmp_path), fmt="png", settings=settings)
    assert dlg2.panels["png"].fields["dpi"].value() == 300
    fmt, path, opts = dlg2.values()
    assert fmt == "png" and path.endswith("x.png") and opts["dpi"] == 300


def test_export_restricted_document_blocked(window, ctrl, tmp_path, monkeypatch, messages):
    path = tmp_path / "vietato.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "Riservato", fontsize=12)
    doc.save(str(path), encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="capo", user_pw="",
             permissions=pymupdf.PDF_PERM_PRINT)
    doc.close()
    window.open_path(str(path))
    called = []
    monkeypatch.setattr(ExportDialog, "exec", lambda d: called.append(1) or 0)
    assert ctrl.export("docx") is None
    assert not called and "copia del contenuto" in messages["error"][0][1]


def test_export_unopened_file_with_password(window, ctrl, encrypted_pdf, tmp_path, monkeypatch, messages):
    from PySide6.QtWidgets import QInputDialog
    monkeypatch.setattr(window.ctx, "ask_open_paths", lambda *a, **k: [encrypted_pdf])
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("segreta", True))
    _accept_export(monkeypatch, tmp_path, "txt")
    result = ctrl.export()
    assert "Riservato" in open(result.paths[0], encoding="utf-8").read()


def test_ocr_never_warning_before_start(window, ctrl, tmp_path, monkeypatch, messages):
    src = pymupdf.open()
    sp = src.new_page()
    sp.insert_text((60, 200), "SCANSIONE", fontsize=40)
    pix = sp.get_pixmap(dpi=100)
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_image(p.rect, stream=pix.tobytes("png"))
    path = tmp_path / "scan.pdf"
    doc.save(str(path))
    window.open_path(str(path))
    shown = []

    def fake_box_exec(box):
        shown.append(box.text())
        box_buttons = {b.text(): b for b in box.buttons()}
        if "Annulla" in box_buttons and "Attiva OCR" in box_buttons:
            box_buttons["Annulla"].click()
        return 0
    monkeypatch.setattr(QMessageBox, "exec", fake_box_exec)

    def fake_exec(dlg):
        dlg.format_box.setCurrentIndex(dlg.format_box.findData("docx"))
        dlg.ocr_box.setCurrentIndex(dlg.ocr_box.findData("never"))
        dlg.folder_edit.setText(str(tmp_path / "o"))
        dlg.validate()
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ExportDialog, "exec", fake_exec)
    assert ctrl.export() is None
    assert shown and "OCR" in shown[0] or "scansionate" in shown[0]
    assert not (tmp_path / "o").exists()


# ------------------------------------------------------------- creazione
def _png(path, size=(400, 300), mode="RGB"):
    from PIL import Image
    Image.new(mode, size, (10, 120, 200) if mode == "RGB" else (10, 120, 200, 128)).save(path)
    return str(path)


def test_create_from_images(window, ctrl, tmp_path, monkeypatch, messages):
    paths = [_png(tmp_path / f"img{n}.png") for n in (10, 2, 1)]
    monkeypatch.setattr(ImagesDialog, "exec", lambda d: (d.sort_box.setCurrentIndex(0), d._sort(),
                                                         QDialog.DialogCode.Accepted)[-1])
    view = ctrl.create_from_images(paths)
    assert view is not None and view.document.page_count == 3
    assert view.document.is_modified and view.document.path is None
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert sorted((round(d[0].rect.width), round(d[0].rect.height))) == [595, 842]   # A4 predefinito


def test_images_dialog_sort_and_options(window, tmp_path):
    paths = [_png(tmp_path / f"img{n}.png") for n in (10, 2)]
    dlg = ImagesDialog(window, paths)
    dlg.sort_box.setCurrentIndex(dlg.sort_box.findData("name"))
    dlg._sort()
    assert [os.path.basename(p) for p in dlg.paths()] == ["img2.png", "img10.png"]
    dlg.list.item(0).setSelected(True)
    dlg._rotate(90)
    opts = dlg.options()
    assert opts["rotations"] == {dlg.paths()[0]: 90} and opts["page_size"] == "A4"


def test_create_blank_and_close_asks(window, ctrl, monkeypatch, messages):
    monkeypatch.setattr(window.ctx, "paper_size", lambda: "A4")
    monkeypatch.setattr(window.ctx, "measure_unit", lambda: "mm")

    def fake_exec(d):
        d.pages.setValue(3)
        d.paper_box.setCurrentIndex(d.paper_box.findData("grid"))
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(BlankDialog, "exec", fake_exec)
    view = ctrl.create_blank()
    assert view.document.page_count == 3
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert abs(d[0].rect.width - 595) <= 1 and abs(d[0].rect.height - 842) <= 1
        assert d[0].get_text().strip() == ""
    asked = []

    def box_exec(box):
        asked.append(box.text())
        next(b for b in box.buttons() if b.text() == "Annulla").click()
        return 0
    monkeypatch.setattr(QMessageBox, "exec", box_exec)
    assert not window.close_view(view)
    assert asked


def test_blank_dialog_rejects_invalid_size(window):
    dlg = BlankDialog(window)
    dlg.size_box.setCurrentIndex(dlg.size_box.findData("custom"))
    dlg.w.setValue(0.5)
    dlg.unit.setCurrentIndex(dlg.unit.findData("pt"))
    dlg.accept()
    assert dlg.result() != QDialog.DialogCode.Accepted
    assert dlg.error.isVisibleTo(dlg) and dlg.error.text()


def test_clipboard(window, ctrl, messages):
    from PySide6.QtGui import QGuiApplication, QImage
    cb = QGuiApplication.clipboard()
    cb.clear()
    ctrl.actions["clipboard"].setEnabled(ctrl.clipboard_supported())
    assert not ctrl.actions["clipboard"].isEnabled()
    img = QImage(1920, 1080, QImage.Format.Format_RGB32)
    img.fill(0x336699)
    cb.setImage(img)
    assert ctrl.clipboard_supported()
    view = ctrl.create_from_clipboard()
    assert view is not None and view.document.page_count == 1
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        info = d[0].get_images(full=True)[0]
        assert (info[2], info[3]) == (1920, 1080)
        assert d.xref_get_key(info[0], "Filter")[1] != "/DCTDecode"
    cb.clear()


def test_create_from_text_file(window, ctrl, tmp_path, messages):
    path = tmp_path / "note.txt"
    path.write_text("Orologio Verdi\nseconda riga", encoding="utf-8")
    views = ctrl.create_from_files([str(path)])
    assert len(views) == 1
    assert "Orologio Verdi" in views[0].document.page_text(0)


# ------------------------------------------------------------- compressione
@pytest.fixture
def photo_pdf(tmp_path):
    import random
    from PIL import Image
    random.seed(1)
    img = Image.effect_noise((1200, 900), 60).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 60), "Foto del quadrante", fontsize=12)
    p.insert_image(pymupdf.Rect(72, 80, 360, 296), stream=buf.getvalue())
    path = tmp_path / "foto.pdf"
    doc.save(str(path))
    return str(path)


def test_compress_document_undoable(window, ctrl, photo_pdf, monkeypatch, messages):
    view = window.open_path(photo_pdf)
    before_file = open(photo_pdf, "rb").read()
    seen = {}

    def fake_exec(d):
        seen["estimates"] = [d.table.item(r, 1).text() for r in range(d.table.rowCount())]
        d.table.selectRow(3)
        d.update_preview()
        seen["preview"] = d.after_label.pixmap() is not None and not d.after_label.pixmap().isNull()
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(CompressDialog, "exec", fake_exec)
    res = ctrl.compress()
    assert res["reduced"] and res["after"] < res["before"]
    assert all(e != "—" for e in seen["estimates"]) and seen["preview"]
    assert view.document.is_modified and "Foto del quadrante" in view.document.page_text(0)
    assert open(photo_pdf, "rb").read() == before_file
    assert view.document.undo()
    assert "Riduzione" in messages["notify"][0][0]


def test_compress_no_reduction(window, ctrl, tmp_path, monkeypatch, messages):
    from foliant.modules.convert import optimize
    path = tmp_path / "piccolo.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "x")
    doc.save(str(path), garbage=4, deflate=True)
    view = window.open_path(str(path))
    monkeypatch.setattr(CompressDialog, "exec", lambda d: QDialog.DialogCode.Accepted)
    # più grande di qualche KB: i byte del documento aperto e la dimensione del file differiscono di
    # qualche byte (secondo elemento di /ID rigenerato a ogni scrittura, a volte come stringa letterale)
    monkeypatch.setattr(optimize, "compress_pdf_bytes", lambda data, *a, **k: data + b"\n%" + b" " * 4096)
    res = ctrl.compress()
    assert res["reduced"] is False and not view.document.is_modified
    assert "Nessuna riduzione" in messages["notify"][0][0]


# ------------------------------------------------------------- riparazione
def test_repair_ui(window, ctrl, sample_pdf, tmp_path, monkeypatch, messages):
    data = bytearray(open(sample_pdf, "rb").read())
    i = data.rfind(b"startxref")
    data[i + 10:i + 14] = b"9999"
    bad = tmp_path / "rotto.pdf"
    bad.write_bytes(bytes(data))
    monkeypatch.setattr(window.ctx, "ask_open_paths", lambda *a, **k: [str(bad)])
    results = ctrl.repair()
    assert len(results) == 1
    out = results[0][1]
    assert out.endswith("rotto_riparato.pdf") and pymupdf.open(out).page_count == 3


def test_repair_jpeg_proposes_creation(window, ctrl, tmp_path, monkeypatch, messages):
    from PIL import Image
    fake = tmp_path / "foto.pdf"
    Image.new("RGB", (50, 40), (200, 0, 0)).save(fake, "JPEG")
    monkeypatch.setattr(window.ctx, "ask_open_paths", lambda *a, **k: [str(fake)])
    created = []
    monkeypatch.setattr(ctrl, "create_from_images", lambda paths=None: created.append(paths))
    ctrl.repair()
    assert "JPEG" in messages["confirm"][0][1] and created


# ------------------------------------------------------------- in blocco
def test_batch_convert_logic(tmp_path, encrypted_pdf):
    from tests.conftest import make_pdf
    from foliant.modules.convert.batch import ERROR, OK, SKIPPED, batch_convert, collect_pdfs
    root = tmp_path / "in"
    for sub in ("a", "b", "c"):
        (root / sub).mkdir(parents=True)
        make_pdf(root / sub / f"{sub}.pdf", pages=1)
    (root / "rotto.pdf").write_bytes(b"non un pdf")
    (root / "note.txt").write_text("x")
    files, rejected = collect_pdfs([str(root), str(root / "a" / "a.pdf"), encrypted_pdf])
    assert len(files) == 5 and rejected == [] and files.count(str(root / "a" / "a.pdf")) == 1
    out = tmp_path / "out"
    items = batch_convert(files, "txt", out_dir=str(out), keep_structure=True, root=str(root))
    status = {os.path.basename(i.path): i.status for i in items}
    assert status["rotto.pdf"] == ERROR and status["protetto.pdf"] == SKIPPED
    assert all(status[f"{s}.pdf"] == OK for s in "abc")
    for sub in "abc":
        assert (out / sub / f"{sub}.txt").exists()


def test_batch_convert_ui(window, ctrl, sample_pdf, tmp_path, monkeypatch, messages):
    from foliant.modules.convert import batch_dialog
    from foliant.modules.convert.batch_dialog import BatchDialog

    def fake_exec(d):
        d.add_paths([sample_pdf])
        d.format_box.setCurrentIndex(d.format_box.findData("txt"))
        d.dest_box.setCurrentIndex(1)
        d.folder_edit.setText(str(tmp_path / "dest"))
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(BatchDialog, "exec", fake_exec)
    reports = []
    monkeypatch.setattr(batch_dialog, "show_batch_report", lambda c, items: reports.append(items))
    monkeypatch.setattr("foliant.modules.convert.batch_dialog.show_batch_report",
                        lambda c, items: reports.append(items))
    items = ctrl.batch_convert()
    assert items and items[0].status == "completato"
    assert (tmp_path / "dest" / "esempio.txt").exists()


def test_translations_cover_ui_strings():
    """Ogni testo italiano dell'interfaccia del modulo ha la traduzione inglese."""
    import ast
    import pathlib
    from foliant.modules.convert.i18n_en import EN
    folder = pathlib.Path(__file__).resolve().parents[3] / "foliant" / "modules" / "convert"
    missing = set()
    for f in folder.glob("*.py"):
        for node in ast.walk(ast.parse(f.read_text(encoding="utf-8"))):
            if isinstance(node, ast.Call) and getattr(node.func, "id", None) == "_" and node.args \
                    and isinstance(node.args[0], ast.Constant) and isinstance(node.args[0].value, str):
                if node.args[0].value not in EN:
                    missing.add(node.args[0].value)
    assert not missing, sorted(missing)[:20]


def test_flatten_ui_undoable(window, ctrl, tmp_path, monkeypatch, messages):
    from foliant.modules.convert.optimize_dialogs import FlattenDialog
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 100), "Testo", fontsize=12)
    p.add_highlight_annot(pymupdf.Rect(70, 88, 120, 104))
    path = tmp_path / "ann.pdf"
    doc.save(str(path))
    view = window.open_path(str(path))
    monkeypatch.setattr(FlattenDialog, "exec", lambda d: QDialog.DialogCode.Accepted)
    ctrl.flatten()
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert list(d[0].annots()) == []
    assert view.document.undo()
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert len(list(d[0].annots())) == 1
    assert len(list(pymupdf.open(str(path))[0].annots())) == 1


def test_scanner_ui_flatbed_two_pages_insert(window, ctrl, sample_pdf, monkeypatch, messages):
    from PIL import Image
    from foliant.modules.convert.create_dialogs import ScannerDialog

    class Fake:
        def list_devices(self):
            return [("sim:0", "Simulato")]

        def scan(self, device, settings, progress=None):
            yield Image.new("RGB", (int(8.27 * settings.dpi), int(11.69 * settings.dpi)), "white")

    ctrl.scanner_backend = Fake()
    view = window.open_path(sample_pdf)

    def fake_exec(d):
        d.dpi.setCurrentIndex(d.dpi.findData(75))
        d.insert.setChecked(True)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ScannerDialog, "exec", fake_exec)
    clicks = {"n": 0}

    def box_exec(box):
        clicks["n"] += 1
        label = "Acquisisci un'altra pagina" if clicks["n"] == 1 else "Fine"
        for b in box.buttons():
            if b.text() == label:
                b.click()
        return 0
    monkeypatch.setattr(QMessageBox, "exec", box_exec)
    ctrl.create_from_scanner()
    assert view.document.page_count == 5
    assert "Pagina 1" in view.document.page_text(0) and "Pagina 2" in view.document.page_text(3)


def test_scanner_none_found(window, ctrl, messages):
    class Empty:
        def list_devices(self):
            return []
    ctrl.scanner_backend = Empty()
    assert ctrl.create_from_scanner() is None
    assert "Nessuno scanner" in messages["info"][0][1]


def test_batch_parallel(tmp_path):
    from tests.conftest import make_pdf
    from foliant.modules.convert.batch import OK, batch_convert
    files = [make_pdf(tmp_path / f"f{k}.pdf", pages=2, text=f"File {k} pagina {{n}}") for k in range(4)]
    (tmp_path / "rotto.pdf").write_bytes(b"%PDF-1.4 rotto")
    files.append(str(tmp_path / "rotto.pdf"))
    items = batch_convert(files, "txt", out_dir=str(tmp_path / "out"), workers=2)
    assert [i.status for i in items[:4]] == [OK] * 4 and items[4].status == "errore"
    for k in range(4):
        assert "File %d pagina 2" % k in (tmp_path / "out" / f"f{k}.txt").read_text(encoding="utf-8")


def test_screen_capture(window, ctrl, monkeypatch, messages):
    from PySide6.QtCore import QRect
    from PySide6.QtGui import QPixmap, QScreen
    from foliant.modules.convert.create_dialogs import CaptureDialog
    pm = QPixmap(800, 600)
    pm.fill()
    monkeypatch.setattr(QScreen, "grabWindow", lambda self, *a: pm)
    steps = iter([CaptureDialog.ANOTHER, QDialog.DialogCode.Accepted])

    def fake_exec(d):
        code = next(steps)
        if code == CaptureDialog.ANOTHER:
            d.preview.set_selection(QRect(100, 50, 300, 200))     # area ritagliata
        return code
    monkeypatch.setattr(CaptureDialog, "exec", fake_exec)
    view = ctrl.create_from_screen(delay_s=0)
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert d.page_count == 2
        sizes = [(img[2], img[3]) for p in d for img in p.get_images(full=True)]
        assert sizes == [(300, 200), (800, 600)]
        info = d[1].get_images(full=True)[0]
        assert d.xref_get_key(info[0], "Filter")[1] == "/FlateDecode"
    monkeypatch.setattr(QScreen, "grabWindow", lambda self, *a: QPixmap())
    n = len(window.views())
    assert ctrl.create_from_screen(delay_s=0) is None
    assert len(window.views()) == n and messages["error"]


def test_photos_dialog_manual_corners(window, ctrl, tmp_path, monkeypatch, messages):
    pytest.importorskip("cv2")
    from PIL import Image
    from foliant.modules.convert.create_dialogs import PhotosDialog
    path = tmp_path / "foto.png"
    Image.new("RGB", (400, 300), "white").save(path)

    def fake_exec(d):
        d.editor._corners[:] = [[10, 10], [390, 20], [380, 290], [20, 280]]
        d.size.setCurrentIndex(d.size.findData(None))
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(PhotosDialog, "exec", fake_exec)
    view = ctrl.create_from_photos([str(path)])
    assert view is not None and view.document.page_count == 1


def test_remove_conformance_ui(window, ctrl, sample_pdf, tmp_path, messages):
    if find_tool("gs") is None:
        pytest.skip("Ghostscript non installato")
    from foliant.modules.convert.engine import convert_pdf
    out = convert_pdf(sample_pdf, tmp_path / "a.pdf", "pdfa-2b")[0]
    view = window.open_path(out)
    assert ctrl.remove_conformance() is True
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert "pdfaid" not in (d.get_xml_metadata() or "")
    assert view.document.undo()


def test_batch_same_names(tmp_path):
    from tests.conftest import make_pdf
    from foliant.modules.convert.batch import batch_convert
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    files = [make_pdf(tmp_path / "a" / "doc.pdf", text="A {n}"), make_pdf(tmp_path / "b" / "doc.pdf", text="B {n}")]
    items = batch_convert(files, "txt", out_dir=str(tmp_path / "out"), policy="overwrite")
    outs = sorted(o for i in items for o in i.outputs)
    assert len(outs) == 2 and len(set(outs)) == 2


def test_export_presets(window, settings, tmp_path):
    dlg = ExportDialog(window, page_count=3, name="x", folder=str(tmp_path), fmt="png", settings=settings)
    dlg.panels["png"].fields["dpi"].setValue(600)
    dlg.save_preset("Stampa")
    dlg2 = ExportDialog(window, page_count=3, name="x", folder=str(tmp_path), fmt="txt", settings=settings)
    assert dlg2.preset_box.findData("Stampa") >= 0
    dlg2.apply_preset("Stampa")
    assert dlg2.fmt == "png" and dlg2.panels["png"].fields["dpi"].value() == 600


def test_epub_metadata_options(sample_pdf, tmp_path):
    import zipfile
    from foliant.modules.convert.engine import convert_pdf
    out = convert_pdf(sample_pdf, tmp_path / "m.epub", "epub", title="Catalogo 2026", author="Verdi", lang="en")[0]
    opf = zipfile.ZipFile(out).read("OEBPS/content.opf").decode()
    assert "<dc:title>Catalogo 2026</dc:title>" in opf and "<dc:creator>Verdi</dc:creator>" in opf
    assert "<dc:language>en</dc:language>" in opf


def test_batch_asks_password_for_locked(window, ctrl, encrypted_pdf, sample_pdf, tmp_path, monkeypatch, messages):
    from PySide6.QtWidgets import QInputDialog
    from foliant.modules.convert import batch_dialog
    from foliant.modules.convert.batch_dialog import BatchDialog

    def fake_exec(d):
        d.add_paths([sample_pdf, encrypted_pdf])
        d.format_box.setCurrentIndex(d.format_box.findData("txt"))
        d.dest_box.setCurrentIndex(1)
        d.folder_edit.setText(str(tmp_path / "dest"))
        d.workers.setValue(1)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(BatchDialog, "exec", fake_exec)
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("segreta", True))
    monkeypatch.setattr(batch_dialog, "show_batch_report", lambda c, items: None)
    items = ctrl.batch_convert()
    assert [i.status for i in items] == ["completato", "completato"]
    assert (tmp_path / "dest" / "protetto.txt").exists()


def test_clipboard_after_windows_destroyed(qtbot, settings):
    """Il segnale globale degli appunti non deve raggiungere azioni di finestre già distrutte."""
    import gc
    from PySide6.QtGui import QGuiApplication
    from PySide6.QtWidgets import QApplication
    from foliant.ui.main_window import MainWindow
    for _k in range(2):
        win = MainWindow(settings)
        win.show()
        win.close()
        win.deleteLater()
        QApplication.processEvents()
        QApplication.sendPostedEvents(None, 0)
        del win
        gc.collect()
    errors = []
    import sys
    old = sys.excepthook
    sys.excepthook = lambda *a: errors.append(a)
    try:
        QGuiApplication.clipboard().setText("x")
        QApplication.processEvents()
    finally:
        sys.excepthook = old
    assert not errors


def test_compress_open_encrypted_keeps_protection(window, ctrl, tmp_path, monkeypatch, messages):
    """Sicurezza: la copia compressa di un PDF aperto con la sola password utente resta cifrata
    con gli stessi permessi (nessuna restrizione aggirata)."""
    import io as _io
    from PIL import Image
    from PySide6.QtWidgets import QInputDialog
    img = Image.effect_noise((900, 700), 60).convert("RGB")
    buf = _io.BytesIO()
    img.save(buf, "PNG")
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 60), "Riservato", fontsize=12)
    p.insert_image(pymupdf.Rect(72, 80, 360, 300), stream=buf.getvalue())
    path = tmp_path / "cifrato.pdf"
    perms = pymupdf.PDF_PERM_PRINT
    doc.save(str(path), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="utente", owner_pw="capo",
             permissions=perms)
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("utente", True))
    window.open_path(str(path))

    def fake_exec(d):
        d.target.setCurrentIndex(d.target.findData("file"))
        d.table.selectRow(3)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(CompressDialog, "exec", fake_exec)
    out = tmp_path / "compresso.pdf"
    monkeypatch.setattr(window.ctx, "ask_save_path", lambda *a, **k: str(out))
    res = ctrl.compress()
    assert res and res.get("output") == str(out)
    result = pymupdf.open(str(out))
    assert result.needs_pass
    assert result.authenticate("utente")
    assert "AES" in result.metadata["encryption"] and "256" in result.metadata["encryption"]
    assert not result.permissions & pymupdf.PDF_PERM_COPY
    assert not result.permissions & pymupdf.PDF_PERM_MODIFY
    assert result.permissions & pymupdf.PDF_PERM_PRINT


def test_create_portfolio_ui(window, ctrl, sample_pdf, tmp_path, monkeypatch, messages):
    from foliant.modules.convert.portfolio_dialog import COL_CUSTOM, COL_DESC, COL_FOLDER, PortfolioDialog
    other = tmp_path / "nota.txt"
    other.write_text("ciao", encoding="utf-8")

    def fake_exec(d):
        d.title.setText("Pratica 42")
        d.table.item(0, COL_FOLDER).setText("Contratti")
        d.table.item(1, COL_FOLDER).setText("Note")
        d.table.item(1, COL_DESC).setText("Appunto del cliente")
        d.custom_name.setText("Cliente")
        d.table.item(0, COL_CUSTOM).setText("Rossi")
        d.table.item(1, COL_CUSTOM).setText("Bianchi")
        d.view.setCurrentIndex(d.view.findData("thumbnails"))
        d.sort.setCurrentIndex(d.sort.findData("custom"))
        d.initial.setCurrentIndex(d.initial.findData(1))
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(PortfolioDialog, "exec", fake_exec)
    view = ctrl.create_portfolio([sample_pdf, str(other)])
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert d.embfile_count() == 2 and "Pratica 42" in d[0].get_text() and "Appunto del cliente" in d[0].get_text()
        cx = int(d.xref_get_key(d.pdf_catalog(), "Collection")[1].split()[0])
        assert d.xref_get_key(cx, "View")[1] == "/T"
        assert d.xref_get_key(cx, "Sort/S")[1] == "/Cliente"
        assert d.xref_get_key(cx, "Schema/Cliente/Subtype")[1] == "/S"
        assert d.xref_get_key(cx, "D")[1] == "<2>nota.txt"
        assert sorted(d.embfile_names()) == ["<1>esempio.pdf", "<2>nota.txt"]
    assert view.document.display_name == "Pratica 42.pdf"


def test_images_as_image_size_ignores_margin(window, ctrl, tmp_path, monkeypatch, messages):
    """«Come l'immagine»: il campo Margini è disattivato e non deve aggiungere una cornice."""
    from PIL import Image
    path = tmp_path / "foto.png"
    Image.new("RGB", (384, 288), (20, 60, 200)).save(path, dpi=(72, 72))

    def fake_exec(d):
        d.size_box.setCurrentIndex(d.size_box.findData(None))
        assert not d.margin.isEnabled()
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ImagesDialog, "exec", fake_exec)
    view = ctrl.create_from_images([str(path)])
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert abs(d[0].rect.width - 384) < 0.5 and abs(d[0].rect.height - 288) < 0.5


def test_compress_document_reduction_survives_save(window, ctrl, tmp_path, monkeypatch, messages):
    """«Applica al documento aperto»: la riduzione annunciata è quella del file salvato."""
    import os
    doc = pymupdf.open()
    for i in range(40):
        p = doc.new_page()
        for k in range(30):
            p.insert_text((50, 50 + k * 20), f"Pagina {i} riga {k} del catalogo orologi usati", fontsize=10)
    path = tmp_path / "quaranta.pdf"
    doc.save(str(path), garbage=0, deflate=True)       # senza flussi di oggetti
    view = window.open_path(str(path))

    def fake_exec(d):
        d.table.selectRow(0)                            # senza perdita
        d.target.setCurrentIndex(d.target.findData("document"))
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(CompressDialog, "exec", fake_exec)
    res = ctrl.compress()
    before = os.path.getsize(path)
    assert res["reduced"], messages["notify"]
    if res["reduced"]:
        assert window.save_view(view)
        saved = os.path.getsize(path)
        # «Dopo» e il file salvato sono due scritture: il secondo elemento di /ID, rigenerato a ogni
        # scrittura, può essere scritto come stringa letterale e cambiare la lunghezza di qualche byte
        assert abs(saved - res["after"]) <= 64 and saved < before
    else:
        assert not view.document.is_modified and os.path.getsize(path) == before


def test_export_uses_thumbnail_selection(window, ctrl, twenty_pages, tmp_path, monkeypatch, messages):
    view = window.open_path(twenty_pages)

    previous = window.selected_pages_provider
    window.ctx.set_selected_pages_provider(lambda v: [1, 4, 6] if v is view else [])
    monkeypatch.setattr(window, "selected_pages_provider", window.selected_pages_provider)
    _ = previous

    def fake_exec(dlg):
        assert dlg.selected == [1, 4, 6]
        dlg.folder_edit.setText(str(tmp_path))
        dlg.format_box.setCurrentIndex(dlg.format_box.findData("txt"))
        dlg.range_edit.setText("selezionate")
        dlg.validate()
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ExportDialog, "exec", fake_exec)
    result = ctrl.export()
    text = open(result.paths[0], encoding="utf-8").read()
    nums = [int(line.split()[-1]) for line in text.replace("\f", "\n").splitlines() if line.startswith("Pagina")]
    assert nums == [2, 5, 7]


def test_repaired_document_banner_and_copy(window, ctrl, sample_pdf, tmp_path, monkeypatch, messages):
    data = bytearray(open(sample_pdf, "rb").read())
    i = data.rfind(b"startxref")
    data[i + 10:i + 14] = b"9999"
    bad = tmp_path / "rotto.pdf"
    bad.write_bytes(bytes(data))
    view = window.open_path(str(bad))
    assert view.document.repaired
    assert "riparato" in window.ctx.banners(view)
    out = tmp_path / "rotto_riparato.pdf"
    monkeypatch.setattr(window.ctx, "ask_save_path", lambda *a, **k: str(out))
    assert ctrl.save_repaired_copy(view) == str(out)
    fixed = pymupdf.open(str(out))
    assert fixed.page_count == 3 and not getattr(fixed, "is_repaired", False)
    assert "riparato" not in window.ctx.banners(view)
    assert open(bad, "rb").read() == bytes(data)


def test_export_runs_non_modal(window, ctrl, sample_pdf, tmp_path, monkeypatch, messages):
    """Le conversioni lunghe non bloccano la finestra: run_task con modal=False e on_done."""
    window.open_path(sample_pdf)
    seen = {}
    original = window.ctx.run_task

    def spy(title, fn, *args, **kwargs):
        if kwargs.get("modal") is False:
            seen["modal"] = False
            seen["on_done"] = kwargs.get("on_done") is not None
        return original(title, fn, *args, **kwargs)
    monkeypatch.setattr(window.ctx, "run_task", spy)
    _accept_export(monkeypatch, tmp_path, "txt")
    result = ctrl.export()
    assert seen == {"modal": False, "on_done": True}
    assert result is not None and os.path.exists(result.paths[0])



def test_conversion_result_is_a_notification(window, ctrl, sample_pdf, tmp_path, monkeypatch, messages):
    """Fine conversione: avviso non bloccante con «Apri» e «Mostra cartella», nessuna finestra modale."""
    window.open_path(sample_pdf)
    opened = []
    monkeypatch.setattr(QMessageBox, "exec", lambda self: opened.append(self.text()) or 0)
    _accept_export(monkeypatch, tmp_path, "txt")
    ctrl.export()
    assert not opened
    text, actions = messages["notify"][-1]
    assert "esempio.txt" in text and actions[:2] == ["Apri", "Mostra cartella"]


def test_export_allowed_after_owner_unlock(window, ctrl, tmp_path, monkeypatch, messages):
    """Dopo lo sblocco con la password dei permessi la conversione è consentita."""
    path = tmp_path / "vietato.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "Riservato", fontsize=12)
    doc.save(str(path), encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="capo", user_pw="",
             permissions=pymupdf.PDF_PERM_PRINT)
    doc.close()
    view = window.open_path(str(path))
    assert not view.document.allows(pymupdf.PDF_PERM_COPY)
    assert view.document.unlock("capo")
    assert view.document.allows(pymupdf.PDF_PERM_COPY)
    _accept_export(monkeypatch, tmp_path / "out", "txt")
    result = ctrl.export()
    assert result is not None and "Riservato" in open(result.paths[0], encoding="utf-8").read()


@pytest.mark.parametrize("policy", ["overwrite", "rename"])
def test_batch_password_second_pass_keeps_unique_names(window, ctrl, tmp_path, monkeypatch, messages, policy):
    """Seconda passata sui file protetti: le uscite pianificate per tutto il lotto vengono riusate."""
    from PySide6.QtWidgets import QInputDialog
    from foliant.modules.convert import batch_dialog
    from foliant.modules.convert.batch_dialog import BatchDialog
    (tmp_path / "in" / "a").mkdir(parents=True)
    (tmp_path / "in" / "b").mkdir(parents=True)
    a = pymupdf.open()
    a.new_page().insert_text((72, 100), "DOCUMENTO A", fontsize=12)
    a.save(str(tmp_path / "in" / "a" / "doc.pdf"))
    b = pymupdf.open()
    b.new_page().insert_text((72, 100), "DOCUMENTO B", fontsize=12)
    b.save(str(tmp_path / "in" / "b" / "doc.pdf"), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="u",
           owner_pw="o")
    out = tmp_path / "out"

    def fake_exec(d):
        d.add_paths([str(tmp_path / "in")])
        d.format_box.setCurrentIndex(d.format_box.findData("txt"))
        d.dest_box.setCurrentIndex(1)
        d.folder_edit.setText(str(out))
        d.keep.setChecked(False)
        d.policy_box.setCurrentIndex(d.policy_box.findData(policy))
        d.workers.setValue(1)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(BatchDialog, "exec", fake_exec)
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("u", True))
    monkeypatch.setattr(batch_dialog, "show_batch_report", lambda c, items: None)
    items = ctrl.batch_convert()
    assert [i.status for i in items] == ["completato", "completato"]
    outputs = [i.outputs[0] for i in items]
    assert len(set(outputs)) == 2
    contents = {os.path.basename(p): open(p, encoding="utf-8").read() for p in outputs}
    assert "DOCUMENTO A" in contents["doc.txt"] and "DOCUMENTO B" in contents["doc (2).txt"]
    assert sorted(os.listdir(out)) == ["doc (2).txt", "doc.txt"]


def test_scanner_dialog_presets_save_export_import(window, settings, tmp_path, monkeypatch):
    """CNV-24: preset predefiniti, salvati, esportati e importati."""
    from PIL import Image
    from PySide6.QtWidgets import QInputDialog

    from foliant.modules.convert.create_dialogs import ScannerDialog
    asked = []
    dlg = ScannerDialog(window, [("sim:0", "Simulato")], settings=settings,
                        preview=lambda dev, s: Image.new("RGB", (80, 110), "white"))
    i = dlg.preset.findText("Documenti in bianco e nero (fatture, contratti)")
    assert i > 0
    dlg.preset.setCurrentIndex(i)
    dlg.preset.activated.emit(i)
    assert dlg.mode.currentData() == "bw" and dlg.ocr.isChecked() and dlg.source.currentData() == "adf"
    assert not dlg.delete_preset.isEnabled()
    dlg.dpi.setCurrentIndex(dlg.dpi.findData(400))
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("Fatture fornitori", True))
    dlg.save_preset.click()
    assert dlg.preset.currentText() == "Fatture fornitori" and dlg.delete_preset.isEnabled()
    out = tmp_path / "preset.json"
    # riservatezza: dialoghi del nucleo (senza cronologia dei file), mai QFileDialog
    monkeypatch.setattr(window.ctx, "ask_save_path", lambda *a, **k: asked.append(a) or str(out))
    dlg.export_preset.click()
    assert asked and asked[0][1] == "Fatture fornitori.json"
    dlg.delete_preset.click()
    assert dlg.preset.findText("Fatture fornitori") < 0
    # nuova finestra: il preset importato torna disponibile con i suoi valori
    dlg2 = ScannerDialog(window, [("sim:0", "Simulato")], settings=settings)
    monkeypatch.setattr(window.ctx, "ask_open_paths", lambda *a, **k: [str(out)])
    dlg2.import_preset.click()
    assert dlg2.preset.currentText() == "Fatture fornitori"
    assert dlg2.dpi.currentData() == 400 and dlg2.mode.currentData() == "bw"
    assert dlg2.settings().dpi == 400
    dlg3 = ScannerDialog(window, [("sim:0", "Simulato")], settings=settings)
    assert dlg3.preset.findText("Fatture fornitori") > 0
    # file non valido: avviso, nessun preset nuovo
    bad = tmp_path / "rotto.json"
    bad.write_text("non json")
    warned = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: warned.append(a[2]))
    monkeypatch.setattr(window.ctx, "ask_open_paths", lambda *a, **k: [str(bad)])
    n = dlg3.preset.count()
    dlg3.import_preset.click()
    assert warned and dlg3.preset.count() == n
    # anteprima
    assert dlg.preview_button.isVisible() or not dlg.isVisible()
    dlg.preview_button.click()
    assert dlg.preview_label.pixmap() is not None and not dlg.preview_label.pixmap().isNull()


@pytest.mark.skipif(find_tool("soffice") is None, reason="LibreOffice non installato")
def test_create_from_office_with_options(window, ctrl, settings, tmp_path, monkeypatch, messages):
    """CNV-18: opzioni dei documenti d'ufficio (fogli di calcolo) ricordate; conferma per i fogli lunghi."""
    import openpyxl

    from foliant.modules.convert import create
    from foliant.modules.convert.create_dialogs import OfficeOptionsDialog
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 30):
        for c in range(1, 26):
            ws.cell(r, c, f"Ref{r}-{c}")
    wb.create_sheet("Altro")["A1"] = "foglio da escludere"
    path = tmp_path / "listino.xlsx"
    wb.save(path)
    seen = {}

    def fake_exec(dlg):
        seen["sheet"] = dlg.sheet_box.isVisibleTo(dlg)
        seen["slides"] = dlg.slide_box.isVisibleTo(dlg)
        dlg.fit_width.setChecked(True)
        dlg.sheets.setCurrentIndex(dlg.sheets.findData("active"))
        dlg._accept()
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(OfficeOptionsDialog, "exec", fake_exec)
    views = ctrl.create_from_files([str(path)])
    assert seen == {"sheet": True, "slides": False}
    assert len(views) == 1 and views[0].document.page_count == 1
    assert "foglio da escludere" not in views[0].document.page_text(0)
    remembered = OfficeOptionsDialog(window, kinds={"sheet"}, settings=window.ctx.settings())
    assert remembered.fit_width.isChecked() and remembered.sheets.currentData() == "active"
    # foglio molto lungo: conferma; se rifiutata non si crea nulla
    monkeypatch.setattr(create, "SHEET_PAGE_LIMIT", 1)
    monkeypatch.setattr(window.ctx, "confirm", lambda t, m: messages["confirm"].append(m) or False)
    assert ctrl.create_from_files([str(path)]) == []
    assert "pagine" in messages["confirm"][-1]
    # annullando le opzioni non si crea nulla
    monkeypatch.setattr(OfficeOptionsDialog, "exec", lambda d: QDialog.DialogCode.Rejected)
    assert ctrl.create_from_files([str(path)]) == []


def test_export_more_files_opens_tools_batch(window, ctrl, sample_pdf, monkeypatch, messages):
    """CNV-17: «Esporta come › Più file…» apre l'elaborazione in blocco (BAT-06) con il formato
    scelto; i formati che BAT-06 non produce restano nella finestra del modulo."""
    from foliant.modules.convert.batch_dialog import BatchDialog
    window.open_path(sample_pdf)

    def pick(fmt):
        def fake_exec(dlg):
            dlg.format_box.setCurrentIndex(dlg.format_box.findData(fmt))
            return ExportDialog.BATCH
        monkeypatch.setattr(ExportDialog, "exec", fake_exec)
    if getattr(window.ctx, "tools_module", None) is None:
        pytest.skip("modulo Strumenti non caricato")
    pick("xlsx")
    batch = ctrl.export()
    converts = [s for s in batch.steps if s.action == "convert"]
    assert len(converts) == 1 and converts[0].merged_params()["format"] == "xlsx"
    pick("docx")
    batch = ctrl.export()
    converts = [s for s in batch.steps if s.action == "convert"]
    assert len(converts) == 1 and converts[0].merged_params()["format"] == "docx"
    seen = {}

    def batch_exec(dlg):
        seen["fmt"] = dlg.format_box.currentData()
        return QDialog.DialogCode.Rejected
    monkeypatch.setattr(BatchDialog, "exec", batch_exec)
    pick("epub")
    assert ctrl.export() is None and seen["fmt"] == "epub"


def test_batch_fifty_files_two_damaged_one_encrypted(tmp_path, encrypted_pdf):
    """CNV-17: 50 PDF, 2 irreparabili e 1 cifrato saltato → 47 DOCX; il report elenca i 2 errori
    con la causa e il file saltato; nessun temporaneo nella cartella di uscita."""
    import docx
    import shutil

    from tests.conftest import make_pdf
    from foliant.modules.convert.batch import ERROR, OK, SKIPPED, batch_convert, report_csv
    src = tmp_path / "in"
    src.mkdir()
    files = [make_pdf(src / f"doc{k:02}.pdf", pages=1) for k in range(47)]
    for name in ("rotto1.pdf", "rotto2.pdf"):
        (src / name).write_bytes(b"%PDF-1.7\n garbage \x00\xff" * 3)
        files.append(str(src / name))
    files.append(shutil.copy(encrypted_pdf, src / "cifrato.pdf"))
    out = tmp_path / "out"
    items = batch_convert(files, "docx", out_dir=str(out), workers=1)
    by = {os.path.basename(i.path): i for i in items}
    assert sum(1 for i in items if i.status == OK) == 47
    assert by["rotto1.pdf"].status == ERROR and by["rotto1.pdf"].message
    assert by["cifrato.pdf"].status == SKIPPED
    produced = sorted(os.listdir(out))
    assert len(produced) == 47 and all(p.endswith(".docx") for p in produced)
    docx.Document(str(out / produced[0]))      # file completi e apribili
    report = report_csv(items)
    assert "rotto1.pdf" in report and "rotto2.pdf" in report and "cifrato.pdf" in report
    assert by["rotto2.pdf"].message in report


def test_local_paper_size_and_unit(window, ctrl, tmp_path, monkeypatch, messages):
    """NFR-26: con formato carta Letter e pollici i nuovi documenti e i dialoghi li propongono."""
    from PIL import Image
    monkeypatch.setattr(window.ctx, "paper_size", lambda: "Letter")
    monkeypatch.setattr(window.ctx, "measure_unit", lambda: "in")
    seen = {}

    def blank_exec(d):
        seen["unit"], seen["w"], seen["step"] = d.unit.currentData(), d.w.value(), d.step.suffix()
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(BlankDialog, "exec", blank_exec)
    view = ctrl.create_blank()
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert (round(d[0].rect.width), round(d[0].rect.height)) == (612, 792)
    assert seen == {"unit": "in", "w": 8.5, "step": " in"}
    img = tmp_path / "foto.png"
    Image.new("RGB", (300, 200), "white").save(img)

    def images_exec(d):
        seen["size"], seen["margin"] = d.size_box.currentData(), d.margin.suffix()
        seen["margin_pt"] = d.options()["margin_pt"]
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ImagesDialog, "exec", images_exec)
    view = ctrl.create_from_images([str(img)])
    assert seen["size"] == "Letter" and seen["margin"] == " in" and abs(seen["margin_pt"] - 28.8) < 0.01
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert (round(d[0].rect.width), round(d[0].rect.height)) in ((612, 792), (792, 612))
    txt = tmp_path / "note.txt"
    txt.write_text("Ricevuta di consegna", encoding="utf-8")
    view = ctrl.create_from_files([str(txt)])[0]
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert (round(d[0].rect.width), round(d[0].rect.height)) == (612, 792)


def _camera_frame():
    import cv2
    import numpy as np

    from tests.modules.convert.test_photo import _sheet
    sheet = _sheet()
    h, w = sheet.shape[:2]
    corners = np.array([[380, 220], [1240, 300], [1300, 1560], [300, 1500]], dtype="float32")
    m = cv2.getPerspectiveTransform(np.array([[0, 0], [w - 1, 0], [w - 1, h - 1], [0, h - 1]], dtype="float32"),
                                    corners)
    bg = np.full((1800, 1600, 3), (60, 90, 70), np.uint8)
    mask = cv2.warpPerspective(np.full((h, w), 255, np.uint8), m, (1600, 1800))
    bg[mask > 0] = cv2.warpPerspective(sheet, m, (1600, 1800))[mask > 0]
    return bg


def test_create_from_camera_two_shots(window, ctrl, monkeypatch, messages):
    """CNV-26: due scatti dalla webcam (simulata) → foto con angoli rilevati → PDF di 2 pagine A4."""
    from foliant.modules.convert.create_dialogs import CameraDialog, PhotosDialog
    frame = _camera_frame()
    log = []

    class FakeCamera:
        def list_devices(self):
            return [(0, "Webcam simulata")]

        def open(self, index):
            log.append(("open", index))

        def read(self):
            return frame

        def close(self):
            log.append("close")

    ctrl.camera_backend = FakeCamera()
    monkeypatch.setattr(window.ctx, "paper_size", lambda: "A4")

    def cam_exec(dlg):
        assert not dlg.done_button.isEnabled()
        dlg.shoot()
        dlg.shoot()
        assert "2" in dlg.status.text() and dlg.done_button.isEnabled()
        dlg.done(QDialog.DialogCode.Accepted)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(CameraDialog, "exec", cam_exec)
    seen = {}

    def photos_exec(dlg):
        seen["found"] = [it["found"] for it in dlg.items]
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(PhotosDialog, "exec", photos_exec)
    view = ctrl.create_from_camera()
    assert seen["found"] == [True, True]
    assert log[0] == ("open", 0) and log[-1] == "close"
    with pymupdf.open(stream=view.document.to_bytes()) as d:
        assert d.page_count == 2 and abs(d[0].rect.width - 595) < 2


def test_create_from_camera_without_device_offers_import(window, ctrl, monkeypatch, messages):
    class NoCamera:
        def list_devices(self):
            return []
    ctrl.camera_backend = NoCamera()
    called = []
    monkeypatch.setattr(ctrl, "create_from_photos", lambda paths=None: called.append(paths) or "importa")
    assert ctrl.create_from_camera() == "importa" and called == [None]
    assert "fotocamera" in messages["confirm"][0][1].lower()


def test_camera_open_error_disables_shot(window):
    from foliant.modules.convert.camera import CameraError
    from foliant.modules.convert.create_dialogs import CameraDialog

    class Busy:
        def open(self, index):
            raise CameraError("occupata")

        def read(self):
            return None

        def close(self):
            pass
    dlg = CameraDialog(window, Busy(), [(0, "Webcam")])
    assert not dlg.shot_button.isEnabled() and "occupata" in dlg.view.text()
    dlg.reject()


def test_cv_camera_lists_no_devices_here():
    from foliant.modules.convert.camera import CvCamera
    cam = CvCamera()
    assert isinstance(cam.list_devices(), list)
    assert cam.read() is None
    cam.close()


def test_photos_ocr_option(window, ctrl, tmp_path, monkeypatch, messages):
    from foliant.modules.convert.create_dialogs import PhotosDialog
    from PIL import Image
    path = tmp_path / "garanzia.png"
    Image.fromarray(_camera_frame()).save(path)
    ocr = []
    monkeypatch.setattr(ctrl, "_with_ocr", lambda data: ocr.append(len(data)) or data)

    def photos_exec(dlg):
        dlg.ocr.setChecked(True)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(PhotosDialog, "exec", photos_exec)
    assert ctrl.create_from_photos([str(path)]) is not None and len(ocr) == 1


def test_batch_dialog_empty_placeholder_and_drop(window, sample_pdf, tmp_path):
    """Collaudo C12: l'elenco vuoto spiega cosa fare; si possono trascinare file e cartelle."""
    from PySide6.QtCore import QMimeData, QPoint, QUrl
    from PySide6.QtGui import QDropEvent
    from PySide6.QtCore import Qt

    from foliant.modules.convert.batch_dialog import BatchDialog
    dlg = BatchDialog(window, window.ctx)
    assert dlg.list.count() == 0 and "trascinali" in dlg.list.placeholder
    dlg.resize(640, 600)
    img = dlg.list.grab().toImage()          # il segnaposto viene disegnato senza errori
    assert not img.isNull()
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(sample_pdf)])
    ev = QDropEvent(QPoint(10, 10), Qt.DropAction.CopyAction, mime, Qt.MouseButton.LeftButton,
                    Qt.KeyboardModifier.NoModifier)
    dlg.list.dropEvent(ev)
    assert dlg.files() == [sample_pdf] and dlg.ok_button.isEnabled()


def test_export_dialog_browse_uses_core_dialog(window, tmp_path, monkeypatch):
    """«Sfoglia…» della destinazione: dialogo del nucleo, nessuna cartella ricordata da Qt."""
    from foliant.modules.convert.export_dialog import ExportDialog
    asked = []
    monkeypatch.setattr(window.ctx, "ask_directory", lambda *a, **k: asked.append(a) or str(tmp_path / "scelta"))
    dlg = ExportDialog(window, page_count=3, name="x", folder=str(tmp_path))
    dlg._browse()
    assert asked and dlg.folder_edit.text() == str(tmp_path / "scelta")
    dlg.close()


def test_convert_keeps_no_paths_in_settings(window, settings, tmp_path):
    """Le preferenze di Converti salvate non contengono la cartella né il nome del documento."""
    from foliant.modules.convert.export_dialog import ExportDialog
    folder = tmp_path / "Pratiche riservate"
    folder.mkdir()
    dlg = ExportDialog(window, page_count=3, name="perizia-rossi", folder=str(folder), fmt="png",
                       settings=settings)
    dlg.save_preset("Mio")
    fmt, path, options = dlg.values()
    assert "perizia-rossi" in path                      # il percorso serve solo per questa esportazione
    dlg._save_settings(fmt, options)
    settings.sync()
    dump = open(settings.fileName(), "rb").read()
    assert b"perizia-rossi" not in dump and b"Pratiche riservate" not in dump
    dlg.close()


def test_export_settings_do_not_keep_document_title(window, settings, tmp_path):
    """Riservatezza: titolo e autore EPUB del singolo documento non restano nelle impostazioni."""
    dlg = ExportDialog(window, page_count=2, name="contratto_rossi", folder=str(tmp_path), fmt="epub",
                       settings=settings)
    dlg.panels["epub"].fields["title"].setText("Contratto Rossi")
    dlg.panels["epub"].fields["author"].setText("Mario Rossi")
    dlg.folder_edit.setText(str(tmp_path))
    assert dlg.validate()
    dlg.accept()
    dlg.save_preset("Libri")
    settings.sync()
    dump = " ".join(str(settings.value(k)) for k in settings.allKeys())
    assert "Rossi" not in dump
