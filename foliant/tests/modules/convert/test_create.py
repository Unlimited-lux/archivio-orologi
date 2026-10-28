"""Test della creazione di PDF (CNV-18…CNV-23, CNV-28)."""

from __future__ import annotations

import io
import re
import shutil
import time
import unicodedata

import pymupdf
import pytest
from PIL import Image

from foliant.modules.convert import create
from foliant.modules.convert.common import ConversionError

HAS_SOFFICE = create.soffice_path() is not None
needs_soffice = pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice non installato")


def _open(data: bytes) -> pymupdf.Document:
    return pymupdf.open("pdf", data)


def _jpeg(path, size, color=(200, 30, 30), exif_orientation=None, block=True):
    im = Image.new("RGB", size, color)
    if block:   # quadrato blu in alto a sinistra per riconoscere il verso
        for x in range(min(size) // 4):
            for y in range(min(size) // 4):
                im.putpixel((x, y), (0, 0, 255))
    kw = {}
    if exif_orientation:
        ex = Image.Exif()
        ex[0x0112] = exif_orientation
        kw["exif"] = ex
    im.save(path, "JPEG", quality=90, **kw)
    return str(path)


def _norm(s: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s))


# ===================================================================== CNV-23
def test_page_sizes_and_units():
    assert set(["A0", "A6", "B4", "B5", "Letter", "Legal", "Tabloid"]) <= set(create.PAGE_SIZES)
    w, h = create.PAGE_SIZES["A4"]
    assert abs(w - 595) < 1 and abs(h - 842) < 1
    assert create.size_from_units(210, 297, "mm") == pytest.approx((595.28, 841.89), abs=0.05)
    assert create.size_from_units(8.5, 11, "in") == (612, 792)
    assert create.size_from_units(100, 200, "pt") == (100, 200)


def test_blank_grid_a4():
    data = create.blank_pdf(size="A4", orientation="portrait", pages=3, paper="grid", step_mm=5)
    pdf = _open(data)
    assert pdf.page_count == 3
    for page in pdf:
        assert abs(page.rect.width - 595) <= 1 and abs(page.rect.height - 842) <= 1
        assert page.get_text().strip() == ""
    # un solo Form XObject condiviso da tutte le pagine
    xobjs = {tuple(x[:1]) for p in pdf for x in p.get_xobjects()}
    assert len(xobjs) == 1
    drawings = pdf[0].get_drawings()
    xs = sorted({round(it[1].x, 3) for d in drawings for it in d["items"]
                 if it[0] == "l" and abs(it[1].x - it[2].x) < 0.01})
    steps = [b - a for a, b in zip(xs, xs[1:])]
    assert len(steps) > 30
    assert all(abs(s - 14.17) <= 0.1 for s in steps)
    ys = sorted({round(it[1].y, 3) for d in drawings for it in d["items"]
                 if it[0] == "l" and abs(it[1].y - it[2].y) < 0.01})
    assert all(abs(b - a - 14.17) <= 0.1 for a, b in zip(ys, ys[1:]))
    # marcato come artefatto di sfondo
    contents = pdf.xref_stream(pdf[0].get_contents()[0])
    assert b"/Artifact" in contents and b"/Background" in contents


def test_blank_user_unit_6000x1000mm():
    size = create.size_from_units(6000, 1000, "mm")
    data = create.blank_pdf(size=size, orientation="landscape", paper="lines")
    pdf = _open(data)
    page = pdf[0]
    assert page.mediabox.width <= 14400 and page.mediabox.height <= 14400
    w, h = create.effective_page_size(page)
    assert abs(w / create.MM - 6000) <= 1 and abs(h / create.MM - 1000) <= 1


@pytest.mark.parametrize("size", [(2, 100), (100, 1e12), (float("nan"), 50)])
def test_blank_bad_size(size):
    with pytest.raises(ConversionError, match="pt"):
        create.blank_pdf(size=size)


def test_blank_1000_pages_fast():
    t = time.perf_counter()
    data = create.blank_pdf(size="A4", pages=1000, paper="dots", background=(1, 1, 0.9))
    assert time.perf_counter() - t < 2
    assert _open(data).page_count == 1000


def test_blank_plain():
    pdf = _open(create.blank_pdf(size="Letter", orientation="landscape", pages=2))
    assert pdf.page_count == 2 and pdf[0].rect.width == pytest.approx(792)


# ===================================================================== CNV-19
def test_natural_sort():
    assert create.sort_images(["/x/img10.png", "/x/img2.png", "/x/Img1.png"], "name") == [
        "/x/Img1.png", "/x/img2.png", "/x/img10.png"]
    assert create.sort_images(["b", "a"], "none") == ["b", "a"]


def test_three_jpegs_auto_orientation(tmp_path):
    paths = [_jpeg(tmp_path / "land.jpg", (400, 200)), _jpeg(tmp_path / "port.jpg", (200, 400)),
             _jpeg(tmp_path / "sq.jpg", (300, 300))]
    margin = 36
    data, warns = create.images_to_pdf(paths, page_size="A4", orientation="auto", margin_pt=margin)
    assert not warns
    pdf = _open(data)
    assert pdf.page_count == 3
    a4w, a4h = create.PAGE_SIZES["A4"]
    assert pdf[0].rect.width == pytest.approx(a4h, abs=1) and pdf[0].rect.height == pytest.approx(a4w, abs=1)
    for i in (1, 2):
        assert pdf[i].rect.width == pytest.approx(a4w, abs=1)
    for page in pdf:
        inner = page.rect + (margin, margin, -margin, -margin)
        bbox = page.get_image_info()[0]["bbox"]
        r = pymupdf.Rect(bbox)
        assert r.x0 >= inner.x0 - 0.5 and r.y0 >= inner.y0 - 0.5
        assert r.x1 <= inner.x1 + 0.5 and r.y1 <= inner.y1 + 0.5
        # centrata
        assert abs((r.x0 + r.x1) / 2 - page.rect.width / 2) < 0.5
        assert abs((r.y0 + r.y1) / 2 - page.rect.height / 2) < 0.5
        # contieni: almeno un lato tocca l'area interna
        assert abs(r.width - inner.width) < 1 or abs(r.height - inner.height) < 1


def test_jpeg_stream_identical(tmp_path):
    path = _jpeg(tmp_path / "foto.jpg", (320, 240))
    data, _w = create.images_to_pdf([path])
    pdf = _open(data)
    xref = pdf[0].get_images()[0][0]
    assert pdf.xref_stream_raw(xref) == open(path, "rb").read()
    assert "DCTDecode" in pdf.xref_get_key(xref, "Filter")[1]
    # "come l'immagine": 72 dpi se assente → 320×240 pt
    assert pdf[0].rect.width == pytest.approx(320, abs=0.5)


def test_jpeg_exif_orientation_6(tmp_path):
    ref = Image.new("RGB", (200, 300), (255, 255, 255))
    for x in range(60):
        for y in range(60):
            ref.putpixel((x, y), (0, 0, 255))
    stored = ref.transpose(Image.Transpose.ROTATE_90)       # 300×200 come lo salverebbe la fotocamera
    ex = Image.Exif()
    ex[0x0112] = 6
    path = tmp_path / "ruotata.jpg"
    stored.save(path, "JPEG", exif=ex, quality=95)
    data, _w = create.images_to_pdf([str(path)])
    pdf = _open(data)
    page = pdf[0]
    assert page.rect.width == pytest.approx(200, abs=0.5) and page.rect.height == pytest.approx(300, abs=0.5)
    # stream ancora identico al file
    xref = page.get_images()[0][0]
    assert pdf.xref_stream_raw(xref) == path.read_bytes()
    pix = page.get_pixmap()
    got = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    diff = sum(1 for x in range(0, 200, 4) for y in range(0, 300, 4)
               if sum(abs(a - b) for a, b in zip(got.getpixel((x, y)), ref.getpixel((x, y)))) > 90)
    assert diff / (50 * 75) <= 0.01


def test_tiff_frames_and_png_alpha(tmp_path):
    frames = [Image.new("RGB", (100, 80), c) for c in ("red", "green", "blue", "yellow")]
    tif = tmp_path / "multi.tiff"
    frames[0].save(tif, save_all=True, append_images=frames[1:])
    png = tmp_path / "alpha.png"
    im = Image.new("RGBA", (64, 64), (255, 0, 0, 0))
    for x in range(32):
        for y in range(64):
            im.putpixel((x, y), (0, 128, 0, 255))
    im.save(png)
    data, warns = create.images_to_pdf([str(tif), str(png)])
    pdf = _open(data)
    assert pdf.page_count == 5
    xref = pdf[4].get_images()[0][0]
    smask = pdf.xref_get_key(xref, "SMask")
    assert smask[0] == "xref"


def test_damaged_and_other_formats(tmp_path):
    bad = tmp_path / "rotta.jpg"
    bad.write_bytes(b"\xff\xd8\xff\xe0 non una vera immagine")
    txt = tmp_path / "nota.png"
    txt.write_bytes(b"ciao")
    gif = tmp_path / "anim.gif"
    g = [Image.new("RGB", (40, 40), c) for c in ("red", "green", "blue")]
    g[0].save(gif, save_all=True, append_images=g[1:])
    bmp = tmp_path / "x.bmp"
    Image.new("RGB", (50, 30), "white").save(bmp)
    webp = tmp_path / "x.webp"
    Image.new("RGB", (50, 30), "white").save(webp)
    data, warns = create.images_to_pdf([str(bad), str(txt), str(gif), str(bmp), str(webp)])
    pdf = _open(data)
    assert pdf.page_count == 3          # GIF: primo fotogramma
    assert len(warns) == 2 and "rotta.jpg" in warns[0]
    data, _w = create.images_to_pdf([str(gif)], gif_all_frames=True)
    assert _open(data).page_count == 3
    with pytest.raises(ConversionError):
        create.images_to_pdf([str(bad)])


def test_dpi_fill_and_recompress(tmp_path):
    path = tmp_path / "dpi.png"
    Image.new("RGB", (300, 150), "white").save(path, dpi=(300, 300))
    data, _w = create.images_to_pdf([str(path)])
    page = _open(data)[0]
    assert page.rect.width == pytest.approx(72, abs=0.5) and page.rect.height == pytest.approx(36, abs=0.5)
    j = _jpeg(tmp_path / "big.jpg", (2000, 1000))
    data, _w = create.images_to_pdf([j], page_size="A6", fit="fill", margin_pt=10, quality="recompress",
                                    max_ppi=100, rotations={j: 90})
    pdf = _open(data)
    page = pdf[0]
    info = page.get_image_info()[0]
    assert max(info["width"], info["height"]) < 2000
    assert page.rect.height > page.rect.width     # ruotata di 90°: verticale


# ===================================================================== CNV-20
MULTI = "Perché così città è già\nΚαλημέρα κόσμε\n你好，世界，中文测试\nمرحبا بالعالم\nשלום עולם\n"


def test_text_multilingual():
    data = create.text_to_pdf(MULTI, header_footer=False)
    pdf = _open(data)
    text = "".join(p.get_text() for p in pdf)
    assert _norm(text) == _norm(MULTI)
    for page in pdf:
        for block in page.get_text("rawdict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    for ch in span["chars"]:
                        assert ch["c"] != "�"


def test_text_wrap_header_footer_and_numbers():
    long = "x" * 1000
    text = "\n".join(f"riga {i}" for i in range(150)) + "\n" + long + "\tfine"
    pdf = _open(create.text_to_pdf(text, title="prova.txt", line_numbers=True))
    assert pdf.page_count >= 3
    all_text = "".join(p.get_text() for p in pdf)
    assert all_text.count("x") >= 1000
    assert "Pagina 1 di" in pdf[0].get_text() and "prova.txt" in pdf[0].get_text()
    for page in pdf:   # nessun testo fuori pagina
        for b in page.get_text("blocks"):
            assert b[2] <= page.rect.width + 0.5


def test_read_text_file_encodings(tmp_path):
    src = "Perché città già più così – àèìòù\n"
    p = tmp_path / "latin1.txt"
    p.write_bytes("Perché città già più così àèìòù\n".encode("iso-8859-1"))
    text, enc = create.read_text_file(p)
    assert text == "Perché città già più così àèìòù\n" and enc in ("iso-8859-1", "cp1252")
    data = create.text_to_pdf(text)
    assert "Perché città già più così àèìòù" in _open(data)[0].get_text()
    p8 = tmp_path / "u8.txt"
    p8.write_bytes(b"\xef\xbb\xbf" + src.encode())
    assert create.read_text_file(p8) == (src, "utf-8-sig")
    p16 = tmp_path / "u16.txt"
    p16.write_bytes(src.encode("utf-16"))
    assert create.read_text_file(p16)[0] == src


MD = """# Primo

Paragrafo con **grassetto**, *corsivo*, ~~barrato~~ e [link](https://example.com).

## Sotto A

- [x] fatto
- [ ] da fare
  - annidato

## Sotto B

```python
print("ciao")
```

| Colonna | Valore |
|:--------|-------:|
| a       | 1      |

![logo](logo.png)

# Secondo

1. uno
2. due

# Terzo

> citazione
"""


def test_markdown_bookmarks(tmp_path):
    Image.new("RGB", (40, 20), "red").save(tmp_path / "logo.png")
    data = create.markdown_to_pdf(MD, base_dir=str(tmp_path))
    pdf = _open(data)
    toc = pdf.get_toc()
    top = [t for t in toc if t[0] == 1]
    assert [t[1] for t in top] == ["Primo", "Secondo", "Terzo"]
    assert [t[1] for t in toc[1:3]] == ["Sotto A", "Sotto B"] and all(t[0] == 2 for t in toc[1:3])
    text = pdf[0].get_text()
    for word in ("grassetto", "barrato", "print(\"ciao\")", "Colonna", "citazione", "annidato"):
        assert word in "".join(p.get_text() for p in pdf)
    links = [lk for p in pdf for lk in p.get_links() if lk["kind"] == pymupdf.LINK_URI]
    assert links and links[0]["uri"] == "https://example.com"
    assert any(p.get_images() for p in pdf)
    fonts = {f[3] for f in pdf[0].get_fonts()}
    assert any("Mono" in f or "Courier" in f for f in fonts), fonts
    assert "grassetto" in text


def test_markdown_html_subset():
    html, _w = create.markdown_to_html("Testo *a* **b** ~~c~~ `d`\n\nSetext\n======\n\n---\n")
    assert "<i>a</i>" in html and "<b>b</b>" in html and "<s>c</s>" in html and "<code>d</code>" in html
    assert '<h1 id="setext">Setext</h1>' in html and "<hr/>" in html


def test_csv_500_rows_header_each_page(tmp_path):
    p = tmp_path / "dati.csv"
    lines = ["Codice;Descrizione;Prezzo"] + [f"C{i:04d};Orologio modello {i};{i},50" for i in range(500)]
    p.write_text("\n".join(lines), encoding="utf-8")
    pdf = _open(create.csv_to_pdf(str(p)))
    assert pdf.page_count > 3
    found = []
    for page in pdf:
        words = page.get_text("words")
        assert words
        # la prima parola del corpo pagina (dopo l'eventuale intestazione) è l'intestazione della tabella
        body = [w for w in words if w[1] > 20 and w[3] < page.rect.height - 25]
        assert min(body, key=lambda w: (round(w[1]), w[0]))[4] == "Codice"
        found += re.findall(r"C\d{4}", page.get_text())
    assert sorted(set(found)) == [f"C{i:04d}" for i in range(500)] and len(found) == 500


def test_csv_wide_landscape_and_text_input():
    text = ",".join(f"colonna_lunga_{i}" for i in range(15)) + "\n" + ",".join("valore" for _ in range(15))
    pdf = _open(create.csv_to_pdf(text))
    assert pdf[0].rect.width > pdf[0].rect.height
    with pytest.raises(ConversionError):
        create.csv_to_pdf("")


# ===================================================================== CNV-21
HTML = """<html><head><title>Prova</title></head><body>
<h1>Uno</h1><p><a href="https://example.com/a">primo</a> e <a href="https://example.org/b">secondo</a>
e <a href="#fondo">vai in fondo</a></p>
<h1>Due</h1>""" + "<p>Testo di riempimento della pagina.</p>" * 120 + """
<h1 id="fondo">Tre</h1><p>Fine.</p></body></html>"""


def test_html_links_and_bookmarks():
    pdf = _open(create.html_to_pdf(HTML))
    links = [(i, lk) for i, p in enumerate(pdf) for lk in p.get_links()]
    uris = sorted(lk["uri"] for _i, lk in links if lk["kind"] == pymupdf.LINK_URI)
    assert uris == ["https://example.com/a", "https://example.org/b"]
    internal = [lk for _i, lk in links if lk["kind"] == pymupdf.LINK_GOTO]
    assert len(internal) == 1
    target = pdf[internal[0]["page"]]
    assert pdf.page_count > 1 and internal[0]["page"] == pdf.page_count - 1
    y = internal[0]["to"].y
    hits = target.search_for("Tre")
    assert hits and any(abs(h.y0 - y) < 30 or abs((target.rect.height - y) - h.y0) < 30 for h in hits)
    toc = pdf.get_toc()
    assert [t[1] for t in toc] == ["Uno", "Due", "Tre"] and all(t[0] == 1 for t in toc)
    assert pdf.metadata["title"] == "Prova"


def test_html_single_page():
    pdf = _open(create.html_to_pdf(HTML, single_page=True))
    assert pdf.page_count == 1
    text = pdf[0].get_text()
    assert "Uno" in text and "Fine." in text and text.count("riempimento") == 120
    big = "<p>riga</p>" * 4000
    pdf = _open(create.html_to_pdf(big, single_page=True))
    page = pdf[0]
    assert pdf.page_count == 1 and page.mediabox.height <= 14400
    assert pdf[0].get_text().count("riga") == 4000
    assert pdf.xref_get_key(page.xref, "UserUnit")[0] in ("real", "int", "float")


def test_url_errors(monkeypatch):
    import urllib.error
    import urllib.request

    def fake(req, timeout=0):
        raise urllib.error.HTTPError(req.full_url, 404, "Not Found", {}, None)
    monkeypatch.setattr(urllib.request, "urlopen", fake)
    with pytest.raises(ConversionError, match="404"):
        create.url_to_html("https://example.com/manca")
    with pytest.raises(ConversionError):
        create.url_to_html("file:///etc/passwd")


def test_url_ok(monkeypatch):
    import urllib.request

    class Resp(io.BytesIO):
        headers = type("H", (), {"get_content_charset": staticmethod(lambda: "utf-8")})()

        def geturl(self):
            return "https://example.com/dir/pagina.html"

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False
    page = "<html><body><a href='altro.html'>x</a><script>alert(1)</script></body></html>"
    monkeypatch.setattr(urllib.request, "urlopen", lambda req, timeout=0: Resp(page.encode()))
    html, final = create.url_to_html("https://example.com/dir/pagina.html")
    assert "https://example.com/dir/altro.html" in html and "alert" not in html


# ===================================================================== CNV-22
def test_clipboard_image_lossless():
    im = Image.new("RGB", (1920, 1080), (10, 20, 30))
    for x in range(0, 1920, 7):
        im.putpixel((x, x % 1080), (255, 255, 255))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    pdf = _open(create.clipboard_to_pdf(image_bytes=buf.getvalue(), text="ignorato"))
    assert pdf.page_count == 1
    info = pdf[0].get_image_info(xrefs=True)[0]
    assert (info["width"], info["height"]) == (1920, 1080)
    assert "DCT" not in pdf.xref_get_key(info["xref"], "Filter")[1]
    pix = pymupdf.Pixmap(pdf, info["xref"])
    assert pix.samples == im.tobytes()
    bmp = io.BytesIO()
    im.save(bmp, "BMP")
    pdf = _open(create.clipboard_to_pdf(image_bytes=bmp.getvalue()))
    assert pymupdf.Pixmap(pdf, pdf[0].get_images()[0][0]).samples == im.tobytes()


def test_clipboard_html_text_and_priority(tmp_path):
    html = ('<html><body><!--StartFragment--><p>Una parola <b>importante</b> e un '
            '<a href="https://example.com/x">collegamento</a>.</p><!--EndFragment--></body></html>')
    pdf = _open(create.clipboard_to_pdf(html=html, text="semplice"))
    spans = [s for b in pdf[0].get_text("dict")["blocks"] for line in b.get("lines", []) for s in line["spans"]]
    bold = [s for s in spans if "importante" in s["text"]]
    assert bold and (bold[0]["flags"] & 16 or "Bold" in bold[0]["font"])
    assert [lk["uri"] for lk in pdf[0].get_links()] == ["https://example.com/x"]
    pdf = _open(create.clipboard_to_pdf(text="Solo testo"))
    assert "Solo testo" in pdf[0].get_text()
    src = pymupdf.open()
    src.new_page().insert_text((72, 72), "copiata")
    pdf_bytes = src.tobytes()
    assert create.clipboard_to_pdf(pdf_bytes=pdf_bytes, html=html) == pdf_bytes
    with pytest.raises(ConversionError):
        create.clipboard_to_pdf()


# ===================================================================== CNV-28
def _batch_files(tmp_path, docx=False):
    files = []
    for i in range(2):
        files.append(_jpeg(tmp_path / f"foto{i}.jpg", (120, 90)))
    png = tmp_path / "schema.png"
    Image.new("RGB", (50, 50), "white").save(png)
    files.append(str(png))
    t = tmp_path / "nota.txt"
    t.write_text("Nota di prova\n", encoding="utf-8")
    files.append(str(t))
    md = tmp_path / "leggimi.md"
    md.write_text("# Titolo\n\ntesto", encoding="utf-8")
    files.append(str(md))
    if docx:
        files.append(_make_docx(tmp_path / "lettera.docx"))
    return files


def test_batch_separate(tmp_path):
    files = _batch_files(tmp_path)
    bad = tmp_path / "archivio.xyz"
    bad.write_bytes(b"???")
    broken = tmp_path / "rotta.png"
    broken.write_bytes(b"no")
    out = tmp_path / "uscita"
    rep = create.batch_create(files + [str(bad), str(broken)], mode="separate", output_dir=str(out))
    assert len(rep["created"]) == len(files)
    assert sorted(p.split("/")[-1] for p in rep["created"]) == sorted(
        f.split("/")[-1].rsplit(".", 1)[0] + ".pdf" for f in files)
    assert [s[0] for s in rep["skipped"]] == [str(bad), str(broken)]
    assert not [p for p in out.iterdir() if p.name.startswith(".")]   # nessun temporaneo rimasto


def test_batch_single(tmp_path):
    files = _batch_files(tmp_path)
    target = tmp_path / "unico.pdf"
    rep = create.batch_create(files, mode="single", output_path=str(target))
    pdf = pymupdf.open(rep["output"])
    toc = pdf.get_toc()
    assert len(toc) == len(files) and all(t[0] == 1 for t in toc)
    assert [t[1] for t in toc] == [f.split("/")[-1].rsplit(".", 1)[0] for f in files]
    assert [t[2] for t in toc] == list(range(1, len(files) + 1))
    assert "Nota di prova" in pdf[3].get_text()
    assert rep["data"] == rep["pdf"]


def test_file_to_pdf_dispatch(tmp_path, sample_pdf):
    data, _w = create.file_to_pdf(sample_pdf)
    assert data == open(sample_pdf, "rb").read()
    with pytest.raises(ConversionError):
        create.file_to_pdf(str(tmp_path / "manca.txt"))
    h = tmp_path / "p.html"
    h.write_text("<h1>Ciao</h1>", encoding="utf-8")
    assert "Ciao" in _open(create.file_to_pdf(str(h))[0])[0].get_text()


# ===================================================================== CNV-18
def _make_docx(path):
    import docx
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = docx.Document()
    for i, title in enumerate(["Introduzione", "Sviluppo", "Conclusione"]):
        d.add_heading(title, level=1)
        p = d.add_paragraph("Paragrafo di testo. ")
        if i < 2:
            url = f"https://example.com/{i}"
            r_id = p.part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), r_id)
            run = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = f"collegamento {i}"
            run.append(t)
            link.append(run)
            p._p.append(link)
    img = io.BytesIO()
    Image.new("RGB", (60, 40), "blue").save(img, "PNG")
    img.seek(0)
    d.add_picture(img)
    d.save(path)
    return str(path)


@needs_soffice
def test_office_docx_and_batch(tmp_path):
    path = _make_docx(tmp_path / "prova.docx")
    pdf = _open(create.office_to_pdf(path))
    assert [t[1] for t in pdf.get_toc() if t[0] == 1] == ["Introduzione", "Sviluppo", "Conclusione"]
    uris = sorted(lk["uri"] for p in pdf for lk in p.get_links() if lk["kind"] == pymupdf.LINK_URI)
    assert uris == ["https://example.com/0", "https://example.com/1"]
    for p in pdf:
        for f in p.get_fonts():
            assert f[1] not in ("", "n/a")    # incorporato (ha un'estensione di file font)
    # XLSX con 3 fogli, nell'ordine
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Primo"
    wb.active["A1"] = "foglio uno"
    for name in ("Secondo", "Terzo"):
        ws = wb.create_sheet(name)
        ws["A1"] = f"foglio {name.lower()}"
    xp = tmp_path / "tre.xlsx"
    wb.save(xp)
    pdf = _open(create.office_to_pdf(str(xp)))
    text = "".join(p.get_text() for p in pdf)
    assert text.index("foglio uno") < text.index("foglio secondo") < text.index("foglio terzo")


@needs_soffice
def test_office_pptx_hidden_slide_and_errors(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    for i in range(10):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Diapositiva {i + 1}"
        if i == 4:
            s._element.set("show", "0")
    pp = tmp_path / "pres.pptx"
    prs.save(pp)
    assert _open(create.office_to_pdf(str(pp))).page_count == 9
    bad = tmp_path / "rotto.docx"
    bad.write_bytes(b"PK\x03\x04 non un documento")
    with pytest.raises(ConversionError):
        create.office_to_pdf(str(bad))


def test_office_missing_soffice(tmp_path, monkeypatch):
    p = tmp_path / "x.docx"
    p.write_bytes(b"PK")
    monkeypatch.setattr(create, "soffice_path", lambda: None)
    with pytest.raises(ConversionError, match="libreoffice-writer"):
        create.office_to_pdf(str(p))


def test_office_cancel(tmp_path):
    from foliant.api import OperationCancelled

    class Cancel:
        def update(self, fraction=None, text=None):
            raise OperationCancelled()
    p = tmp_path / "x.odt"
    p.write_bytes(b"PK")
    if not HAS_SOFFICE:
        pytest.skip("LibreOffice non installato")
    with pytest.raises(OperationCancelled):
        create.office_to_pdf(str(p), progress=Cancel())


def test_cancel_images(tmp_path):
    from foliant.api import OperationCancelled

    class Cancel:
        def update(self, fraction=None, text=None):
            raise OperationCancelled()
    with pytest.raises(OperationCancelled):
        create.images_to_pdf([_jpeg(tmp_path / "a.jpg", (10, 10))], progress=Cancel())


def test_shutil_reference():
    assert HAS_SOFFICE == (shutil.which("soffice") is not None or create.soffice_path() is not None)


def test_add_ocr_layer_keeps_images(tmp_path):
    import io as _io
    from foliant.modules.convert.extract import add_ocr_layer, ocr_available
    if not ocr_available():
        pytest.skip("tesseract non installato")
    src = pymupdf.open()
    sp = src.new_page()
    sp.insert_text((60, 150), "Ricevuta di acquisto", fontsize=30)
    png = sp.get_pixmap(dpi=200).tobytes("png")
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_image(p.rect, stream=png)
    xref = p.get_images()[0][0]
    raw = doc.xref_stream(xref)
    out, pages = add_ocr_layer(doc.tobytes(), "ita")
    res = pymupdf.open(stream=out)
    assert pages == [1] and "Ricevuta" in res[0].get_text()
    imgs = res[0].get_images(full=True)
    assert len(imgs) == 1 and res.xref_stream(imgs[0][0]) == raw
