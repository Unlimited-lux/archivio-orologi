"""AI-25 (generazione di elementi grafici) e AI-27 (conversione assistita dall'IA)."""

from __future__ import annotations

import base64
import json
import threading
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pymupdf
import pytest

from foliant.api import Document
from foliant.modules.ai import convert_assist, graphics
from foliant.modules.ai.config import Profile
from foliant.modules.ai.providers import NETWORK, REFUSED, AIError, OpenAICompatProvider

from .conftest import FakeProvider


# ------------------------------------------------------------------ utilità
def make_png(fg=(200, 20, 20), bg=None, size=64) -> bytes:
    """Quadrato colorato al centro; `bg=None` → sfondo trasparente, altrimenti opaco."""
    data = bytearray()
    for y in range(size):
        for x in range(size):
            inside = size // 4 <= x < 3 * size // 4 and size // 4 <= y < 3 * size // 4
            if inside:
                data += bytes((*fg, 255))
            elif bg is None:
                data += bytes((0, 0, 0, 0))
            else:
                data += bytes((*bg, 255))
    return pymupdf.Pixmap(pymupdf.csRGB, size, size, bytes(data), True).tobytes("png")


VARIANT_COLORS = [(200, 20, 20), (20, 160, 40), (30, 60, 220)]


@pytest.fixture
def image_server(monkeypatch):
    """Endpoint locale compatibile OpenAI: /v1/images/generations con 3 varianti b64."""
    for var in ("http_proxy", "https_proxy", "HTTP_PROXY", "HTTPS_PROXY", "all_proxy", "ALL_PROXY"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("no_proxy", "127.0.0.1,localhost")
    monkeypatch.setenv("NO_PROXY", "127.0.0.1,localhost")
    state = {"requests": [], "mode": "ok"}

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def do_POST(self):
            length = int(self.headers.get("Content-Length", 0))
            body = json.loads(self.rfile.read(length).decode("utf-8"))
            state["requests"].append({"path": self.path, "body": body,
                                      "auth": self.headers.get("Authorization")})
            if state["mode"] == "refuse":
                payload = json.dumps({"error": {"code": "content_policy_violation",
                                                "message": "Your request was rejected by the safety system."}})
                self.send_response(400)
            else:
                data = [{"b64_json": base64.b64encode(make_png(c)).decode("ascii")} for c in VARIANT_COLORS]
                payload = json.dumps({"created": 1, "data": data})
                self.send_response(200)
            raw = payload.encode("utf-8")
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    state["base"] = f"http://127.0.0.1:{server.server_address[1]}/v1"
    yield state
    server.shutdown()
    server.server_close()


def image_profile(base: str, image_model: str = "img-1") -> Profile:
    return Profile.from_catalog("custom", base_url=base, model="chat-1", image_model=image_model,
                                verified=True, consent=True)


def run_now(title, fn):
    return fn()


# ===================================================================== AI-25
def test_generate_three_variants_choose_stamp_and_undo(qtbot, image_server, tmp_path):
    from foliant.modules.ai.graphics_view import GenerateGraphicDialog

    profile = image_profile(image_server["base"])
    provider = OpenAICompatProvider(profile.base_url, "chiave-immagini", profile.model)
    library = graphics.Library(tmp_path / "grafica")
    dlg = GenerateGraphicDialog(profile, provider, run_now, kind="stamp", library=library)
    qtbot.addWidget(dlg)
    assert dlg.generate_button.isEnabled()
    dlg.description.setPlainText("Sigillo rotondo con un orologio")
    dlg.variants_spin.setValue(3)
    dlg.generate_button.click()

    assert dlg.previews.count() == 3
    assert dlg.generate_button.text() == "Rigenera"
    req = image_server["requests"][0]
    assert req["path"] == "/v1/images/generations"
    assert req["body"]["response_format"] == "b64_json"
    assert req["body"]["n"] == 3 and req["body"]["model"] == "img-1"
    assert "Sigillo rotondo" in req["body"]["prompt"]
    assert req["auth"] == "Bearer chiave-immagini"

    dlg.previews.setCurrentRow(1)
    assert dlg.use_button.isEnabled()
    dlg.use_button.click()
    assert dlg.result() == dlg.DialogCode.Accepted
    png = dlg.chosen_png
    assert png == dlg.variants[1]
    items = library.items()
    assert len(items) == 1 and items[0].kind == "stamp" and items[0].png() == png
    meta = json.loads((tmp_path / "grafica" / f"{items[0].id}.json").read_text())
    assert meta["prompt"] == "Sigillo rotondo con un orologio" and meta["created"]

    doc = Document.new()
    rect = pymupdf.Rect(100, 100, 228, 228)
    how = graphics.apply_as_stamp(doc, 0, rect, png)
    assert how == "annot"
    annots = [(a.type[1], a.rect) for a in doc.pdf[0].annots()]
    assert annots and annots[0][0] == "Stamp"
    pix = doc.pdf[0].get_pixmap(clip=rect)
    assert pix.pixel(pix.width // 2, pix.height // 2) == VARIANT_COLORS[1]     # aspetto = variante scelta
    assert pix.pixel(2, 2) == (255, 255, 255)                                 # trasparenza conservata
    assert doc.undo_label == "Timbro generato"
    doc.undo()
    assert list(doc.pdf[0].annots()) == []
    assert not doc.can_undo


def test_refusal_becomes_refused_error(image_server):
    image_server["mode"] = "refuse"
    provider = OpenAICompatProvider(image_server["base"], "k", "chat-1")
    with pytest.raises(AIError) as info:
        graphics.generate_images(provider, "img-1", "qualcosa", n=2)
    assert info.value.kind == REFUSED


def test_profile_without_image_model_disables_generation(qtbot):
    from foliant.modules.ai.graphics_view import GenerateGraphicDialog

    profile = image_profile("http://127.0.0.1:9/v1", image_model="")
    ok, reason = graphics.can_generate(profile)
    assert not ok and "modello per la generazione di immagini" in reason
    anthropic = Profile.from_catalog("anthropic", image_model="x")
    ok2, reason2 = graphics.can_generate(anthropic)
    assert not ok2 and reason2

    dlg = GenerateGraphicDialog(profile, lambda: pytest.fail("nessuna richiesta"), run_now, kind="watermark")
    qtbot.addWidget(dlg)
    assert not dlg.generate_button.isEnabled()
    assert not dlg.disabled_label.isHidden()
    assert dlg.disabled_label.text() == reason
    assert dlg.generate_button.toolTip() == reason


def test_remove_uniform_background_makes_corners_transparent():
    png = make_png(fg=(10, 10, 120), bg=(250, 250, 250))
    assert not graphics.has_transparency(png)
    out = graphics.remove_uniform_background(png, tolerance=20)
    assert graphics.has_transparency(out)
    pix = pymupdf.Pixmap(out)
    for x, y in ((0, 0), (63, 0), (0, 63), (63, 63)):
        assert pix.pixel(x, y)[3] == 0
    assert pix.pixel(32, 32)[3] == 255
    # sfondo non uniforme: invariato
    half = bytearray()
    for _y in range(32):
        for x in range(32):
            half += bytes((255, 255, 255) if x < 16 else (0, 0, 0))
    split = pymupdf.Pixmap(pymupdf.csRGB, 32, 32, bytes(half), False).tobytes("png")
    assert graphics.background_color(split) is None
    assert graphics.remove_uniform_background(split) == split


def test_watermark_background_and_vector_stamp_are_single_undo_steps():
    doc = Document.new(pages=2)
    png = make_png(bg=(240, 240, 200))
    graphics.apply_as_watermark(doc, [0, 1], png, opacity=0.4)
    assert doc.undo_label == "Filigrana generata"
    assert all(doc.pdf[i].get_images() for i in range(2))
    doc.undo()
    assert not any(doc.pdf[i].get_images() for i in range(2))

    graphics.apply_as_background(doc, [1], png)
    assert doc.pdf[1].get_images() and not doc.pdf[0].get_images()
    doc.undo()

    stamp = graphics.vector_text_stamp("Approvato", "#c0392b")
    assert stamp.startswith(b"%PDF")
    graphics.apply_vector_stamp(doc, 0, (300, 80, 500, 160), stamp)
    assert "APPROVATO" in doc.pdf[0].get_text()
    doc.undo()
    assert "APPROVATO" not in doc.pdf[0].get_text()
    assert graphics.vector_stamp_png(stamp, dpi=72).startswith(b"\x89PNG")


# ===================================================================== AI-27
@pytest.fixture
def report_pdf(tmp_path):
    path = tmp_path / "relazione-riservata.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 90), "Relazione annuale", fontsize=22, fontname="hebo")
    page.insert_textbox(pymupdf.Rect(72, 130, 520, 220),
                        "Il laboratorio ha restaurato quaranta orologi da tasca durante l'anno.", fontsize=11)
    page.insert_textbox(pymupdf.Rect(72, 240, 520, 300), "Seconda sezione del testo.", fontsize=11)
    doc.save(str(path))
    doc.close()
    pdf = pymupdf.open(str(path))
    yield pdf, path
    pdf.close()


def heading_classifier(request):
    payload = json.loads(request.messages[0].content)
    blocks = []
    for b in payload["blocks"]:
        blocks.append({"id": b["id"], "role": "heading1" if "Relazione annuale" in b["text"] else "paragraph"})
    return json.dumps({"blocks": blocks, "tables": []})


def _document_xml(path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def _heading_paragraph(xml: str, text: str) -> str:
    import re
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        if text in p:
            return p
    raise AssertionError(f"{text!r} non trovato")


def test_page_blocks_have_position_and_style(report_pdf):
    pdf, _path = report_pdf
    blocks = convert_assist.page_blocks(pdf, 0)
    assert [b["id"] for b in blocks] == list(range(len(blocks)))
    title = blocks[0]
    assert title["text"] == "Relazione annuale" and title["bold"] and title["size"] == 22
    assert set(title) == {"id", "page", "bbox", "text", "font", "size", "bold"}


@pytest.mark.parametrize("use_python_docx", [True, False])
def test_heading1_from_ai_gets_heading1_style(job, fake_provider, report_pdf, tmp_path, monkeypatch,
                                              use_python_docx):
    pdf, path = report_pdf
    if not use_python_docx:
        monkeypatch.setattr(convert_assist, "_docx_available", lambda: False)
    fake_provider.responses = [heading_classifier]
    structure = convert_assist.classify(job, pdf, [0])
    assert structure[0].role(0) == "heading1"
    out = tmp_path / "out.docx"
    report = convert_assist.to_docx(pdf, [0], structure, out)
    assert report["pages_without_ai"] == [] and report["ai_used"]

    para = _heading_paragraph(_document_xml(out), "Relazione annuale")
    assert '<w:pStyle w:val="Heading1"/>' in para
    other = _heading_paragraph(_document_xml(out), "quaranta orologi")
    assert "Heading" not in other
    import docx
    styles = {p.text: p.style.name for p in docx.Document(str(out)).paragraphs}
    assert styles["Relazione annuale"] == "Heading 1"

    # inviati solo i blocchi: testo, posizione e stile; mai il file né il suo nome
    sent = fake_provider.all_text()
    assert "Relazione annuale" in sent and "quaranta orologi" in sent
    assert '"bbox"' in sent and '"font"' in sent and '"size"' in sent and '"bold"' in sent
    assert path.name not in sent and "relazione-riservata" not in sent
    assert "%PDF" not in sent
    req = fake_provider.requests[0]
    assert not any(m.images for m in req.messages)


def test_network_error_still_produces_file_without_ai(job, fake_provider, report_pdf, tmp_path):
    pdf, _path = report_pdf
    fake_provider.responses = [AIError(NETWORK, "unreachable")]
    structure = convert_assist.classify(job, pdf, [0])        # non solleva
    assert dict(structure) == {} and not structure.ai_used
    assert structure.report and "IA non usata" in structure.report[0]
    out = tmp_path / "out.docx"
    report = convert_assist.to_docx(pdf, [0], structure, out)
    assert out.exists() and "Relazione annuale" in _document_xml(out)
    assert report["pages_without_ai"] == [0] and not report["ai_used"]
    assert any("IA non usata" in m for m in report["messages"])
    assert "Heading1" not in _document_xml(out)


@pytest.mark.parametrize("answer", [
    {"blocks": [{"id": 99, "role": "heading1"}], "tables": []},                       # blocco inesistente
    {"blocks": [{"id": 0, "role": "titolone"}], "tables": []},                        # ruolo sconosciuto
    {"blocks": [{"id": 0, "role": "heading1"}],
     "tables": [{"rows": 2, "cols": 2, "cells": [
         {"block": 1, "row": 0, "col": 0, "rowspan": 1, "colspan": 2},
         {"block": 2, "row": 0, "col": 1}]}]},                                        # celle sovrapposte
    "non è JSON",
])
def test_incoherent_answer_ignores_that_page(job, fake_provider, report_pdf, tmp_path, answer):
    pdf, _path = report_pdf
    fake_provider.responses = [answer if isinstance(answer, str) else json.dumps(answer)]
    structure = convert_assist.classify(job, pdf, [0])
    assert 0 not in structure and 0 in structure.ignored and structure.report
    report = convert_assist.to_docx(pdf, [0], structure, tmp_path / "x.docx")
    assert report["pages_without_ai"] == [0]


@pytest.mark.parametrize("use_python_docx", [True, False])
def test_table_with_merged_cells(job, fake_provider, report_pdf, tmp_path, monkeypatch, use_python_docx):
    pdf, _path = report_pdf
    if not use_python_docx:
        monkeypatch.setattr(convert_assist, "_docx_available", lambda: False)
    fake_provider.responses = [json.dumps({
        "blocks": [{"id": 0, "role": "heading2"}, {"id": 1, "role": "cell"}, {"id": 2, "role": "cell"}],
        "tables": [{"rows": 2, "cols": 2, "cells": [
            {"block": 1, "row": 0, "col": 0, "rowspan": 1, "colspan": 2},
            {"block": 2, "row": 1, "col": 1}]}]})]
    structure = convert_assist.classify(job, pdf, [0])
    assert len(structure[0].tables) == 1
    out = tmp_path / "t.docx"
    convert_assist.to_docx(pdf, [0], structure, out)
    import docx
    d = docx.Document(str(out))
    assert d.paragraphs[0].style.name == "Heading 2"
    assert len(d.tables) == 1
    t = d.tables[0]
    assert "quaranta orologi" in t.cell(0, 0).text
    assert t.cell(1, 1).text == "Seconda sezione del testo."
    assert 'w:gridSpan w:val="2"' in _document_xml(out)


def test_pick_graphic_public_api_handles_consent(window):
    from .conftest import FakeProvider, make_active_profile
    services = window.ctx.ai_services
    services.provider_factory = lambda p: FakeProvider()
    opened = []

    def fake_exec(dlg):
        opened.append(dlg.kind)
        dlg.chosen_png = b"\x89PNG-prova"
        return 1
    services.controller.exec_dialog = fake_exec
    assert services.pick_graphic("sticker") is None and not opened         # nessun profilo: niente finestra
    p = make_active_profile(services.store)
    p.consent = False
    services.store.save(p)
    services.panel.ask_consent = lambda profile: False
    assert services.pick_graphic("sticker") is None and not opened         # informativa non accettata
    services.panel.ask_consent = lambda profile: True
    assert services.pick_graphic("sticker") == b"\x89PNG-prova" and opened == ["sticker"]
    assert services.store.active().consent
