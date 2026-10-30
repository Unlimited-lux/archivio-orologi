"""Controlla i risultati della prova dell'installatore. Eseguito dal Foliant installato:
    foliant-cli.exe -I -c "exec(open(r'…\\verifica_risultati.py', encoding='utf-8').read())" CARTELLA

CARTELLA contiene:
    originale.pdf        la scansione di prova, intatta;
    sessione.pdf         lo stesso PDF aperto, modificato e salvato dall'app (sessione.py);
    docx/                OCR in italiano + conversione in Word (foliant-cli blocco);
    pdf/                 OCR in italiano + compressione + PDF/A con Ghostscript (foliant-cli blocco);
    gs-ebook.pdf         compressione diretta con il Ghostscript incluso (-dPDFSETTINGS=/ebook);
    tabella-xlsx/        il PDF con la tabella (crea_pdf_tabella.py) in Excel (foliant-cli blocco);
    tabella-pptx/        lo stesso in PowerPoint (foliant-cli blocco);
    tabella-docx/        lo stesso in Word, testo scorrevole (foliant-cli blocco);
    tabella-fedele/      lo stesso in Word, impaginazione fedele con pdf2docx (converti_fedele.py).
I file Office si aprono con openpyxl, python-pptx e python-docx del Foliant installato: anche
questo prova che le librerie sono nel pacchetto. I nomi effettivi dei file non compaiono nei
messaggi: si stampano solo i controlli.
"""

import datetime as dt
import glob
import os
import sys
import zipfile

import docx
import openpyxl
import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(errors="replace")         # console di Windows in cp1252
base = sys.argv[1]
ATTESE = ("perizia", "orologio", "città", "perché", "garanzia", "ginevra")
# contenuto di crea_pdf_tabella.py, come testo (PowerPoint, Word) e come valori di Excel: codici
# con lo zero iniziale come testo, anni e importi come numeri, date come date
TITOLO = "Inventario degli orologi"
INTESTAZIONE = ("Codice", "Modello", "Città", "Anno", "Prezzo", "Revisione")
TABELLA_TESTO = (
    ("0451", "Speedmaster", "Bienne", "1969", "4.250,00", "12/03/2024"),
    ("0452", "Submariner", "Ginevra", "1972", "7.900,50", "05/11/2023"),
    ("0453", "Calatrava", "Ginevra", "1985", "3.100,75", "21/06/2022"),
    ("0454", "Royal Oak", "Le Brassus", "1978", "18.400,00", "30/09/2021"),
)
TABELLA = (
    ("0451", "Speedmaster", "Bienne", 1969, 4250.0, dt.datetime(2024, 3, 12)),
    ("0452", "Submariner", "Ginevra", 1972, 7900.5, dt.datetime(2023, 11, 5)),
    ("0453", "Calatrava", "Ginevra", 1985, 3100.75, dt.datetime(2022, 6, 21)),
    ("0454", "Royal Oak", "Le Brassus", 1978, 18400.0, dt.datetime(2021, 9, 30)),
)
TESTO_TABELLA = (TITOLO, *INTESTAZIONE, *(v for riga in TABELLA_TESTO for v in riga))
COMMENTO = "Verificare il prezzo del Royal Oak."
errori = []


def controlla(condizione, messaggio):
    print(("ok      " if condizione else "ERRORE  ") + messaggio)
    if not condizione:
        errori.append(messaggio)


def unico(pattern):
    trovati = glob.glob(os.path.join(base, pattern))
    controlla(len(trovati) == 1, f"un solo risultato per {pattern}")
    return trovati[0] if trovati else None


originale = os.path.join(base, "originale.pdf")
size_originale = os.path.getsize(originale)
with pymupdf.open(originale) as doc:
    controlla(not doc[0].get_text().strip(), "la scansione di prova non ha testo")

with pymupdf.open(os.path.join(base, "sessione.pdf")) as doc:
    controlla("Controllato con Foliant" in doc[0].get_text(), "sessione: modifica salvata nel PDF")

percorso = unico("docx/*.docx")
if percorso:
    testo = "\n".join(p.text for p in docx.Document(percorso).paragraphs).lower()
    mancanti = [w for w in ATTESE if w not in testo]
    controlla(not mancanti, f"DOCX: testo italiano riconosciuto (mancano: {mancanti})")

percorso = unico("pdf/*.pdf")
if percorso:
    with pymupdf.open(percorso) as doc:
        testo = doc[0].get_text().lower()
        meta = doc.metadata or {}
        mancanti = [w for w in ATTESE if w not in testo]
        controlla(not mancanti, f"PDF: livello OCR in italiano (mancano: {mancanti})")
        controlla("ita" in (meta.get("keywords") or ""), "PDF: OCR eseguito con la lingua ita")
        controlla("ghostscript" in (meta.get("producer") or "").lower(),
                  f"PDF/A prodotto da Ghostscript ({meta.get('producer')})")
        catalogo = doc.pdf_catalog()
        controlla(doc.xref_get_key(catalogo, "Metadata")[0] == "xref", "PDF/A: metadati XMP presenti")
    size = os.path.getsize(percorso)
    print(f"        dimensione: {size_originale:,} -> {size:,} byte")
    controlla(size < size_originale, "PDF compresso più piccolo dell'originale")

percorso = os.path.join(base, "gs-ebook.pdf")
with pymupdf.open(percorso) as doc:
    controlla(doc.page_count == 1, "Ghostscript incluso: compressione /ebook leggibile")
print(f"        Ghostscript /ebook: {size_originale:,} -> {os.path.getsize(percorso):,} byte")
controlla(os.path.getsize(percorso) < size_originale, "Ghostscript incluso: file più piccolo")


def testo_docx(documento) -> str:
    """Paragrafi e celle delle tabelle (anche annidate) di un documento python-docx."""
    parti = [p.text for p in documento.paragraphs]
    tabelle = list(documento.tables)
    while tabelle:
        tabella = tabelle.pop()
        for riga in tabella.rows:
            for cella in riga.cells:
                parti += [p.text for p in cella.paragraphs]
                tabelle += cella.tables
    return "\n".join(parti)


# PDF con la tabella → Excel: una riga per voce, con i valori tipizzati
percorso = unico("tabella-xlsx/*.xlsx")
if percorso:
    righe = []
    for foglio in openpyxl.load_workbook(percorso).worksheets:
        righe += [tuple(r) for r in foglio.iter_rows(values_only=True) if any(v is not None for v in r)]
    controlla(any(r[:6] == INTESTAZIONE for r in righe), "XLSX: intestazione della tabella")
    mancanti = [voce[0] for voce in TABELLA if voce not in [r[:6] for r in righe]]
    controlla(not mancanti, f"XLSX: codici come testo, anni e importi come numeri, date come date "
                            f"(righe diverse: {mancanti})")

# PDF con la tabella → PowerPoint: una diapositiva con testo, immagine e commento nelle note
percorso = unico("tabella-pptx/*.pptx")
if percorso:
    prs = Presentation(percorso)
    controlla(len(prs.slides) == 1, f"PPTX: una diapositiva per pagina ({len(prs.slides)})")
    controlla(prs.slide_width < prs.slide_height, "PPTX: diapositiva con il formato della pagina (verticale)")
    if len(prs.slides):
        diapositiva = prs.slides[0]
        forme = list(diapositiva.shapes)
        testo = "\n".join(f.text_frame.text for f in forme if f.has_text_frame)
        mancanti = [w for w in TESTO_TABELLA if w not in testo]
        controlla(not mancanti, f"PPTX: testo modificabile della pagina (mancano: {mancanti})")
        controlla(any(f.shape_type == MSO_SHAPE_TYPE.PICTURE for f in forme), "PPTX: immagine della pagina")
        note = diapositiva.notes_slide.notes_text_frame.text if diapositiva.has_notes_slide else ""
        controlla(COMMENTO in note, "PPTX: commento del PDF nelle note della diapositiva")

# PDF con la tabella → Word (testo scorrevole, python-docx): titolo, tabella di Word, commento
percorso = unico("tabella-docx/*.docx")
if percorso:
    documento = docx.Document(percorso)
    controlla(any(p.text == TITOLO and p.style.name.startswith("Heading") for p in documento.paragraphs),
              "DOCX: titolo come stile di titolo")
    celle = [tuple(c.text for c in r.cells) for t in documento.tables for r in t.rows]
    mancanti = [riga[0] for riga in (INTESTAZIONE,) + TABELLA_TESTO if riga not in celle]
    controlla(not mancanti, f"DOCX: tabella di Word con tutte le celle (righe mancanti: {mancanti})")
    with zipfile.ZipFile(percorso) as z:
        commenti = z.read("word/comments.xml").decode("utf-8") if "word/comments.xml" in z.namelist() else ""
    controlla(COMMENTO in commenti, "DOCX: commento del PDF come commento di Word")

# PDF con la tabella → Word con impaginazione fedele (pdf2docx)
percorso = unico("tabella-fedele/*.docx")
if percorso:
    testo = testo_docx(docx.Document(percorso))
    mancanti = [w for w in TESTO_TABELLA if w not in testo]
    controlla(not mancanti, f"DOCX fedele (pdf2docx): testo e tabella (mancano: {mancanti})")

# LibreOffice resta esterno e facoltativo: se manca, l'app spiega come installarlo
from foliant.core import platform
if platform.find_tool("soffice") is None:
    controlla("libreoffice.org" in platform.missing_tool_message("soffice"),
              "LibreOffice assente: il messaggio indica dove scaricarlo")
else:
    print("        LibreOffice presente sulla macchina di prova")

if errori:
    raise SystemExit(f"{len(errori)} controlli non riusciti")
print("risultati: ok")
