"""Prova del pacchetto per Windows (eseguita da foliant-cli.exe in GitHub Actions): strumenti
inclusi trovati nella cartella «tools», lingue OCR presenti, tutti i moduli importabili, librerie e
modelli della conversione in Word, Excel e PowerPoint presenti e funzionanti.

    foliant-cli.exe -I -c "exec(open('packaging/windows/check_bundle.py', encoding='utf-8').read())"

La conversione si prova anche dai sorgenti, su Linux (tests/test_prova_windows.py)."""

import importlib
import os
import subprocess
import tempfile
import zipfile
from pathlib import Path


def check_tools():
    from foliant.core import platform
    tools = {name: platform.find_tool(name) for name in ("tesseract", "gs", "qpdf")}
    print("strumenti:", tools)
    assert all(path and "\\tools\\" in path for path in tools.values()), tools
    langs = subprocess.run([tools["tesseract"], "--list-langs"], capture_output=True, text=True).stdout.split()
    assert all(code in langs for code in ("ita", "eng", "osd")), langs


def check_modules():
    import foliant
    import foliant.modules as modules
    resources = Path(foliant.__file__).resolve().parent / "resources"
    assert (resources / "foliant.svg").is_file(), resources
    icons = len(list((resources / "icons").glob("*.svg")))
    print("icone:", icons)
    assert icons > 100, icons
    found = list(modules.discover())
    print("moduli:", len(found))
    assert len(found) >= 12, [name for _o, name, _m in found]
    broken = [name for _order, name, mod in modules.discover() if isinstance(mod, Exception)]
    print("moduli non importabili:", broken)
    assert not broken, broken


def check_office_conversion():
    """Conversione da PDF a Word, Excel e PowerPoint: le librerie sono importate solo al momento
    della conversione, quindi un errore del pacchetto si vedrebbe solo lì. Si controllano le
    librerie (con quelle di pdf2docx), i file di modello di python-docx e python-pptx e una
    conversione per formato."""
    missing = []
    for name in ("docx", "pptx", "openpyxl", "et_xmlfile", "lxml.etree", "pdf2docx", "fitz", "cv2",
                 "fontTools.ttLib", "PIL.Image"):
        try:
            importlib.import_module(name)
        except Exception as e:  # noqa: BLE001 - si elencano tutte prima di fallire
            missing.append(f"{name} ({e})")
    print("librerie della conversione non importabili:", missing)
    assert not missing, missing
    import docx
    import openpyxl
    import pptx
    import pymupdf

    from foliant.modules.convert.engine import run_conversion
    templates = {docx: ("default.docx", "default-header.xml", "default-footer.xml", "default-comments.xml"),
                 pptx: ("default.pptx", "notesMaster.xml", "notes.xml", "theme.xml")}
    missing = [f"{package.__name__}/templates/{name}" for package, names in templates.items() for name in names
               if not (Path(package.__file__).parent / "templates" / name).is_file()]
    print("modelli di python-docx e python-pptx mancanti:", missing)
    assert not missing, missing

    with tempfile.TemporaryDirectory(prefix="foliant-prova-") as tmp:
        # due pagine con intestazione e piè di pagina ripetuti (sezioni di Word), una tabella con i
        # bordi (foglio di Excel) e un commento (commento di Word, note di PowerPoint)
        pdf = pymupdf.open()
        for n in (1, 2):
            page = pdf.new_page(width=595, height=842)
            page.insert_text((72, 50), "Archivio di prova", fontsize=10)
            page.insert_text((72, 800), f"Pagina {n} di 2", fontsize=10)
            page.insert_text((72, 160), f"Prova del pacchetto, pagina {n}", fontsize=16)
        page = pdf[0]
        rows = (("Codice", "Prezzo"), ("0451", "4.250,00"), ("0452", "7.900,50"))
        for r in range(len(rows) + 1):
            page.draw_line((72, 200 + 24 * r), (312, 200 + 24 * r))
        for x in (72, 192, 312):
            page.draw_line((x, 200), (x, 200 + 24 * len(rows)))
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                page.insert_text((78 + 120 * c, 217 + 24 * r), text, fontsize=10)
        page.add_text_annot((330, 210), "Commento di prova")
        source = os.path.join(tmp, "prova.pdf")
        pdf.save(source)
        pdf.close()

        def convert(name, fmt, **options):
            paths = run_conversion(source, os.path.join(tmp, name), fmt, **options).paths
            assert len(paths) == 1 and os.path.getsize(paths[0]) > 0, (name, paths)
            return paths[0]

        def word_parts(path):
            """Parti XML del documento Word (testo, commenti, intestazioni e piè di pagina)."""
            docx.Document(path)                                         # si apre con python-docx
            with zipfile.ZipFile(path) as z:
                return {name: z.read(name).decode("utf-8") for name in z.namelist() if name.startswith("word/")}

        parts = word_parts(convert("scorrevole.docx", "docx"))
        checks = {"tabella": "4.250,00" in parts["word/document.xml"],
                  "commento": "Commento di prova" in parts.get("word/comments.xml", ""),
                  "intestazione": any(n.startswith("word/header") and "Archivio di prova" in t
                                      for n, t in parts.items()),
                  "piè di pagina": any(n.startswith("word/footer") and "Pagina" in t for n, t in parts.items())}
        print("Word (testo scorrevole):", checks)
        assert all(checks.values()), checks
        body = word_parts(convert("fedele.docx", "docx", layout="faithful"))["word/document.xml"]
        checks = {"testo": "Prova del pacchetto" in body, "tabella": "4.250,00" in body}
        print("Word (impaginazione fedele, pdf2docx):", checks)
        assert all(checks.values()), checks
        values = [c.value for ws in openpyxl.load_workbook(convert("tabella.xlsx", "xlsx")).worksheets
                  for row in ws.iter_rows() for c in row]
        print("Excel:", values)
        assert "0451" in values and 4250 in values and 7900.5 in values, values
        slides = pptx.Presentation(convert("pagine.pptx", "pptx")).slides
        texts = [s.text_frame.text for s in slides[0].shapes if s.has_text_frame]
        notes = slides[0].notes_slide.notes_text_frame.text if slides[0].has_notes_slide else ""
        print("PowerPoint:", len(slides), "diapositive; note:", notes)
        assert len(slides) == 2 and any("Prova del pacchetto" in t for t in texts), texts
        assert "Commento di prova" in notes, notes


if __name__ == "__main__":
    check_tools()
    check_modules()
    check_office_conversion()
    print("pacchetto: ok")
